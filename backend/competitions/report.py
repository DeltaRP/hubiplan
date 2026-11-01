"""Auswertung eines Wettbewerbs als PDF und Excel."""

from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from exchange.export import safe_cell
from finance.report import _styles, _table
from organizations.theme import effective_theme

from . import services
from .models import Competition, Scoring


def filename(comp: Competition) -> str:
    base = re.sub(r"[^A-Za-z0-9ÄÖÜäöüß_-]+", "-", comp.title).strip("-") or "Wettbewerb"
    return f"Auswertung-{base}"


def _meta(comp: Competition) -> list[str]:
    parts = [f"{comp.date:%d.%m.%Y}"]
    if not comp.own and comp.organizer:
        parts.append(f"Ausrichter: {comp.organizer}")
    for v in (comp.occasion, comp.discipline, comp.location):
        if v:
            parts.append(v)
    return parts


def _scoring_info(sc: Scoring) -> str:
    if not sc.competition.own:
        return f"{sc.field_size} Teilnehmer insgesamt" if sc.field_size else ""
    if sc.kind == Scoring.Kind.UNIT:
        mode = {
            "beste": f"Summe der besten {sc.team_size or ''}",
            "summe": "Summe aller",
            "durchschnitt": "Durchschnitt aller",
        }[sc.team_mode]
        return f"{mode} · {sc.prize_places} Preisplätze"
    parts = []
    if sc.shots:
        parts.append(
            f"{sc.shots} Schuss, davon {sc.counted_shots} gewertet"
            if sc.counted_shots and sc.counted_shots != sc.shots
            else f"{sc.shots} Schuss"
        )
    parts.append(f"{sc.prize_places} Preisplätze")
    return " · ".join(parts)


def pdf(comp: Competition) -> bytes:
    if comp.form == Competition.Form.BIRD:
        return bird_pdf(comp)
    org = comp.organization
    data = services.standings(comp)
    scorings = {s.pk: s for s in comp.scorings.all()}
    st = _styles()
    n, r = st["normal"], st["right"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    story: list = [
        Paragraph(f"<b>{escape(comp.title)}</b>", st["title"]),
        Paragraph(escape(" · ".join([org.name, *_meta(comp)])), n),
        Paragraph(
            f"Stand {timezone.localdate():%d.%m.%Y} · {comp.get_status_display()}", st["small"]
        ),
    ]
    for item in comp.details or []:
        story.append(
            Paragraph(
                f"{escape(item.get('label', ''))}: {escape(item.get('value', ''))}", st["small"]
            )
        )
    story.append(Spacer(1, 4 * mm))
    for block in data["scorings"]:
        sc = scorings[block["scoring"]]
        story.append(Paragraph(f"<b>{escape(block['name'])}</b>", st["h2"]))
        story.append(Paragraph(escape(_scoring_info(sc)), st["small"]))
        has_so = any(row.get("shootoff") for row in block["rows"])
        if not comp.own:
            story.append(_foreign_table(block, sc, primary, st))
            continue
        if sc.kind == Scoring.Kind.UNIT:
            head = ["Platz", "Einheit", "Schützen", "Bester Schütze", "Wert"]
            widths = [14 * mm, 52 * mm, 20 * mm, 52 * mm, 24 * mm]
        else:
            head = ["Platz", "Name", "Einheit / Verein", "Wert"]
            widths = [14 * mm, 70 * mm, 54 * mm, 24 * mm]
        if has_so:
            head.append("Stechen")
            widths = [w * 0.9 for w in widths] + [18 * mm]
        rows = [[Paragraph(f"<b>{h}</b>", r if h in ("Wert", "Stechen") else n) for h in head]]
        for row in block["rows"]:
            place = f"{row['place']}." if row.get("place") else "–"
            if row.get("prize"):
                place = f"<b>{place}</b>"
            if sc.kind == Scoring.Kind.UNIT:
                cells = [
                    place,
                    escape(row["name"]),
                    str(row["count"]),
                    escape(row["best"]),
                    escape(row["value_text"] or row.get("reason", "")),
                ]
            else:
                name = escape(row["name"])
                if row.get("out_of_competition"):
                    name += " (a. K.)"
                cells = [place, name, escape(row["unit_name"] or row["club"]), row["value_text"]]
            if has_so:
                cells.append(" / ".join(row.get("shootoff", {}).values()))
            rows.append(
                [
                    Paragraph(c, r if i >= len(cells) - (2 if has_so else 1) else n)
                    for i, c in enumerate(cells)
                ]
            )
        story.append(_table(rows, widths, primary, total_row=False))
        if block["open"]:
            story.append(
                Paragraph(
                    "Stechen offen: "
                    + "; ".join(
                        f"Platz {t['place']}: {', '.join(t['names'])}" for t in block["open"]
                    ),
                    st["small"],
                )
            )
    stats = data["stats"]
    if stats:
        story.append(Paragraph("<b>Kennzahlen</b>", st["h2"]))
        lines = [
            ["Schützen", str(stats["count"])],
            ["davon Mitglieder / Gäste", f"{stats['members']} / {stats['guests']}"],
            ["Durchschnitt", stats["average"]],
            ["Höchster / niedrigster Wert", f"{stats['max']} / {stats['min']}"],
        ]
        story.append(
            _table(
                [[Paragraph(a, n), Paragraph(b, r)] for a, b in lines],
                [70 * mm, 40 * mm],
                primary,
                total_row=False,
            )
        )
        if stats["per_unit"]:
            rows = [
                [
                    Paragraph(f"<b>{h}</b>", n if i == 0 else r)
                    for i, h in enumerate(
                        [
                            "Einheit",
                            "Schützen",
                            "Mitglieder",
                            "Beteiligung",
                            "Durchschnitt",
                            "Bester Wert",
                        ]
                    )
                ]
            ]
            for u in stats["per_unit"]:
                part = f"{u['participation']} %" if u["participation"] is not None else ""
                rows.append(
                    [Paragraph(escape(u["name"]), n)]
                    + [
                        Paragraph(str(x), r)
                        for x in (u["count"], u["members"], part, u["average"], u["best"])
                    ]
                )
            story.append(Spacer(1, 3 * mm))
            story.append(
                _table(
                    rows,
                    [50 * mm, 20 * mm, 22 * mm, 24 * mm, 26 * mm, 24 * mm],
                    primary,
                    total_row=False,
                )
            )

    buf = io.BytesIO()
    title = f"Auswertung {comp.title}"

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(20 * mm, 10 * mm, f"{org.name} · {title}")
        canvas.drawRightString(190 * mm, 10 * mm, f"Seite {doc.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
        author=org.name,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _place_text(row: dict, sc: Scoring) -> str:
    if not row.get("place"):
        return "–"
    text = f"{row['place']}."
    if sc.field_size:
        text += f" von {sc.field_size}"
    return text


def _foreign_table(block: dict, sc: Scoring, primary, st):
    n, r = st["normal"], st["right"]
    team = sc.kind == Scoring.Kind.TEAM
    head = ["Platz", "Mannschaft" if team else "Name", "Schützen" if team else "Einheit", "Wert"]
    rows = [[Paragraph(f"<b>{h}</b>", r if h == "Wert" else n) for h in head]]
    for row in block["rows"]:
        second = ", ".join(row.get("members", [])) if team else (row["unit_name"] or row["club"])
        note = f" ({row['note']})" if row.get("note") else ""
        rows.append(
            [
                Paragraph(escape(_place_text(row, sc)), n),
                Paragraph(escape(row["name"] + note), n),
                Paragraph(escape(second), n),
                Paragraph(escape(row["value_text"]), r),
            ]
        )
    return _table(rows, [24 * mm, 56 * mm, 64 * mm, 22 * mm], primary, total_row=False)


def _bird_numbers(data: dict) -> list[tuple[str, str]]:
    stats = data["stats"]
    lines = [
        ("Schüsse insgesamt", str(stats["shots"] or "")),
        ("Schützen", str(stats["starters"])),
        ("davon Mitglieder / Gäste", f"{stats['members']} / {stats['guests']}"),
    ]
    if stats.get("rounds"):
        lines.append(("Runden", str(stats["rounds"])))
    if stats.get("duration_minutes") is not None:
        h, m = divmod(stats["duration_minutes"], 60)
        lines.append(("Dauer", f"{h} Std. {m} Min." if h else f"{m} Min."))
    for u in stats["per_unit"]:
        lines.append((f"Pfänder {u['name']}", str(u["count"])))
    return lines


def bird_pdf(comp: Competition) -> bytes:
    """Pfänderliste mit Kennzahlen (Vogelschießen)."""
    from . import bird

    org = comp.organization
    data = bird.state(comp)
    st = _styles()
    n, r = st["normal"], st["right"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    story: list = [
        Paragraph(f"<b>{escape(comp.title)}</b>", st["title"]),
        Paragraph(escape(" · ".join([org.name, *_meta(comp)])), n),
        Paragraph(
            f"Stand {timezone.localdate():%d.%m.%Y} · {comp.get_status_display()}", st["small"]
        ),
        Spacer(1, 4 * mm),
        Paragraph("<b>Pfänderliste</b>", st["h2"]),
    ]
    head = ["Pfand", "Schütze", "Einheit / Verein", "Schuss Nr.", "Schüsse", "Preis"]
    rows = [[Paragraph(f"<b>{h}</b>", r if h in ("Schuss Nr.", "Schüsse") else n) for h in head]]
    for p in data["prizes"]:
        rows.append(
            [
                Paragraph(escape(p["name"]), n),
                Paragraph(escape(p["winner_name"] or "–"), n),
                Paragraph(escape(p["unit_name"] or p["club"]), n),
                Paragraph(str(p["shot_number"] or ""), r),
                Paragraph(str(p["shots_for"] or ""), r),
                Paragraph(escape(p["award"]), n),
            ]
        )
    story.append(
        _table(
            rows,
            [30 * mm, 44 * mm, 34 * mm, 18 * mm, 16 * mm, 30 * mm],
            primary,
            total_row=False,
        )
    )
    story.append(Paragraph("<b>Kennzahlen</b>", st["h2"]))
    story.append(
        _table(
            [[Paragraph(escape(a), n), Paragraph(escape(b), r)] for a, b in _bird_numbers(data)],
            [70 * mm, 40 * mm],
            primary,
            total_row=False,
        )
    )
    return _doc(comp, f"Pfänderliste {comp.title}", story)


def lot_list(comp: Competition, blank_rows: int = 10) -> bytes:
    """Losliste fürs Vogelschießen auf Papier: Schützen nach Losnummer mit Feld für Striche
    (gezählte Schüsse) und Pfand, dazu leere Zeilen für Nachzügler."""
    from . import bird

    org = comp.organization
    st = _styles()
    n = st["normal"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    starters = [s for s in comp.starters.select_related("unit") if s.presence != "zurueckgezogen"]
    starters.sort(key=lambda s: (bird.lot(s.number) is None, bird.lot(s.number) or 0, s.name))
    head = ["Los", "Name", "Einheit / Verein", "Schüsse (Striche)", "Pfand"]
    widths = [14 * mm, 58 * mm, 40 * mm, 125 * mm, 20 * mm]
    rows = [[Paragraph(f"<b>{escape(h)}</b>", n) for h in head]]
    for s in starters:
        rows.append(
            [
                Paragraph(escape(s.number), n),
                Paragraph(escape(s.name), n),
                Paragraph(escape(s.unit.name if s.unit_id else s.club), n),
                "",
                "",
            ]
        )
    prizes = ", ".join(p.name for p in comp.prizes.all())

    def story(blank: int) -> list:
        return [
            Paragraph(f"<b>Losliste {escape(comp.title)}</b>", st["title"]),
            Paragraph(escape(" · ".join([org.name, *_meta(comp)])), n),
            Paragraph(escape(f"Pfänder: {prizes}"), st["small"]),
            Spacer(1, 4 * mm),
            _grid(rows + [[""] * len(head) for _ in range(blank)], widths, primary),
        ]

    return _filled(comp, f"Losliste {comp.title}", story, blank_rows, wide=True)


def prize_sheet(comp: Competition) -> bytes:
    """Leere Pfänderliste zum Ausfüllen am Stand."""
    org = comp.organization
    st = _styles()
    n = st["normal"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    head = ["Pfand", "Schütze", "Los", "Schuss Nr.", "Schüsse für das Pfand", "Uhrzeit"]
    widths = [32 * mm, 58 * mm, 14 * mm, 20 * mm, 30 * mm, 18 * mm]
    rows = [[Paragraph(f"<b>{escape(h)}</b>", n) for h in head]]
    for p in comp.prizes.all():
        rows.append([Paragraph(escape(p.name), n)] + [""] * (len(head) - 1))
    story = [
        Paragraph(f"<b>Pfänderliste {escape(comp.title)}</b>", st["title"]),
        Paragraph(escape(" · ".join([org.name, *_meta(comp)])), n),
        Spacer(1, 4 * mm),
        _grid(rows, widths, primary, height=12 * mm),
        Spacer(1, 6 * mm),
        Paragraph("Schüsse insgesamt: ____________", n),
    ]
    return _doc(comp, f"Pfänderliste {comp.title}", story)


def _doc(comp: Competition, title: str, story: list, *, wide: bool = False) -> bytes:
    return _build(comp, title, story, wide=wide)[0]


def _filled(comp: Competition, title: str, make_story, least: int, *, wide: bool = False):
    """PDF mit mindestens ``least`` leeren Zeilen, aufgefüllt bis die letzte Seite voll ist,
    damit keine Seite mit nur einer einzelnen Zeile entsteht (Ralf 2026-11-01).
    ``make_story(blank)`` baut den Inhalt mit ``blank`` leeren Zeilen."""
    body, pages = _build(comp, title, make_story(least), wide=wide)
    if least == 0:  # ausdrücklich ohne leere Zeilen
        return body
    low, high = least, least + 60  # 60 leere Zeilen sind mehr als eine Seite
    while low < high:  # größte Zahl leerer Zeilen, die noch auf gleich viele Seiten passt
        mid = (low + high + 1) // 2
        b, p = _build(comp, title, make_story(mid), wide=wide)
        if p == pages:
            low, body = mid, b
        else:
            high = mid - 1
    if low != least:
        body = _build(comp, title, make_story(low), wide=wide)[0]
    return body


def _build(comp: Competition, title: str, story: list, *, wide: bool = False):
    from reportlab.lib.pagesizes import landscape

    org = comp.organization
    buf = io.BytesIO()
    count = {"pages": 0}

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(20 * mm, 10 * mm, f"{org.name} · {title}")
        canvas.drawRightString((277 if wide else 190) * mm, 10 * mm, f"Seite {doc.page}")
        canvas.restoreState()
        count["pages"] = doc.page

    SimpleDocTemplate(
        buf,
        pagesize=landscape(A4) if wide else A4,
        leftMargin=20 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=title,
        author=org.name,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue(), count["pages"]


def _grid(rows, widths, primary, height=8 * mm):
    from reportlab.platypus import Table

    return Table(
        rows,
        colWidths=widths,
        rowHeights=[None] + [height] * (len(rows) - 1),
        repeatRows=1,
        hAlign="LEFT",
        style=[
            ("LINEBELOW", (0, 0), (-1, 0), 0.8, primary),
            ("GRID", (0, 1), (-1, -1), 0.3, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ],
    )


ORDERS = {"name": "Name", "einheit": "Einheit", "nummer": "Startnummer"}


def shooting_list(comp: Competition, order: str = "name", blank_rows: int = 10, empty=False):
    """Schießliste zum Eintragen auf Papier: Starter mit leeren Feldern je Schuss (oder
    für die Summe) und Unterschrift, dazu leere Zeilen für Nachmeldungen. Nach Einheit
    bekommt jede Einheit eine eigene Tabelle mit Feldern für die Zugwertung (Ralf 2026-11-01)."""
    org = comp.organization
    st = _styles()
    n = st["normal"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    sc = comp.scorings.filter(kind=Scoring.Kind.INDIVIDUAL).order_by("sort_order", "pk").first()
    starters = [] if empty else list(comp.starters.select_related("unit"))
    if order == "einheit":
        starters.sort(key=lambda s: (s.unit.name if s.unit_id else "~", s.name))
    elif order == "nummer":
        starters.sort(
            key=lambda s: (
                not s.number,
                int(s.number) if s.number.isdigit() else 0,
                s.number,
                s.name,
            )
        )
    else:
        starters.sort(key=lambda s: s.name)
    shots = sc.shots if sc and sc.shots and sc.shots <= 10 else 0
    wide = shots > 5
    total = (257 if wide else 172) * mm
    tail = ["gewertet" if shots else (sc.unit_label if sc else "Punkte"), "Aufsicht"]
    tail_widths = [10 * mm] * shots + [20 * mm, 22 * mm]
    shot_heads = [str(i) for i in range(1, shots + 1)]
    # Schon erfasste Werte eindrucken (Ralf 2026-11-01); Einzelschüsse werden nicht erfasst
    values = {}
    if sc and not empty:
        values = {v.starter_id: services.fmt(v.value, sc.decimals) for v in sc.scores.all()}

    def table(people, blank: int, with_unit: bool):
        head = ["Nr.", "Name"] + (["Einheit / Verein"] if with_unit else []) + shot_heads + tail
        widths = [12 * mm, 0] + ([40 * mm] if with_unit else []) + tail_widths
        widths[1] = total - sum(widths)
        rows = [[Paragraph(f"<b>{escape(h)}</b>", n) for h in head]]
        for s in people:
            cells = [Paragraph(escape(s.number), n), Paragraph(escape(s.name), n)]
            if with_unit:
                cells.append(Paragraph(escape(s.unit.name if s.unit_id else s.club), n))
            rest = [""] * (len(head) - len(cells))
            rest[-2] = values.get(s.pk, "")
            rows.append(cells + rest)
        rows += [[""] * len(head) for _ in range(blank)]
        return _grid(rows, widths, primary)

    info = []
    if sc:
        info.append(_scoring_info(sc))

    def heading() -> list:
        return [
            Paragraph(f"<b>Schießliste {escape(comp.title)}</b>", st["title"]),
            Paragraph(escape(" · ".join([org.name, *_meta(comp)])), n),
            Paragraph(
                escape(" · ".join(info + [f"sortiert nach {ORDERS.get(order, 'Name')}"])),
                st["small"],
            ),
            Spacer(1, 4 * mm),
        ]

    from organizations.models import Department

    units = list(Department.objects.filter(organization=org).order_by("sort_order", "name"))
    if order != "einheit" or not units:

        def story(blank: int) -> list:
            return heading() + [table(starters, blank, True)]

        # Leere Liste: genau eine volle Seite; sonst die letzte Seite mit leeren Zeilen füllen
        least = 1 if empty else blank_rows
        return _filled(comp, f"Schießliste {comp.title}", story, least, wide=wide)

    from reportlab.platypus import KeepTogether

    team_scorings = list(comp.scorings.filter(kind=Scoring.Kind.UNIT).order_by("sort_order", "pk"))
    per_unit = 6 if empty else 2  # leere Zeilen je Einheit für Nachmeldungen
    team_rows: dict[tuple[int, int], dict] = {}
    if not empty and team_scorings:
        for block in services.standings(comp)["scorings"]:
            for row in block["rows"] if block["kind"] == Scoring.Kind.UNIT else []:
                team_rows[(block["scoring"], row["unit"])] = row
    raw = {v.starter_id: v.value for v in sc.scores.all()} if sc and not empty else {}
    guests = [s for s in starters if not s.unit_id]

    def story(blank: int) -> list:
        out = heading()
        for unit in units:
            people = [s for s in starters if s.unit_id == unit.pk]
            block = [
                Paragraph(f"<b>{escape(unit.name)}</b>", st["normal"]),
                Spacer(1, 1.5 * mm),
                table(people, per_unit, False),
            ]
            for tsc in team_scorings:
                filled = _team_cells(tsc, team_rows.get((tsc.pk, unit.pk)), raw)
                block += [Spacer(1, 2 * mm), _team_box(tsc, total, primary, n, filled)]
            out.append(KeepTogether(block + [Spacer(1, 6 * mm)]))
        out += [
            Paragraph("<b>Gäste und ohne Einheit</b>", st["normal"]),
            Spacer(1, 1.5 * mm),
            table(guests, blank, True),
        ]
        return out

    least = 1 if empty else blank_rows
    return _filled(comp, f"Schießliste {comp.title}", story, least, wide=wide)


def _team_cells(sc: Scoring, row: dict | None, raw: dict) -> dict[str, str]:
    """Schon feststehende Werte der Zugwertung einer Einheit für die Felder im Kasten."""
    if not row or row.get("value") is None:
        return {}
    src = sc.source
    decimals = src.decimals if src else 0
    counted = [raw[pk] for pk in row.get("counted", []) if raw.get(pk) is not None]
    counted.sort(reverse=src.higher_better if src else True)
    out = {f"{i}.": services.fmt(v, decimals) for i, v in enumerate(counted, 1)}
    if sc.team_mode == "durchschnitt":
        out |= {
            "Summe": services.fmt(sum(counted), decimals),
            "Anzahl": str(len(counted)),
            "Durchschnitt": row["value_text"],
        }
    else:
        out["Summe"] = row["value_text"]
    if row.get("place"):
        out["Platz"] = str(row["place"])
    return out


def _team_box(sc: Scoring, total, primary, n, filled: dict[str, str] | None = None):
    """Felder für die Wertung einer Einheit: je gewertetem Ergebnis eine Zelle, Summe
    (beim Durchschnitt auch Anzahl), Platz."""
    from reportlab.platypus import Table

    if sc.team_mode == "beste" and sc.team_size and sc.team_size <= 10:
        heads = [f"{i}." for i in range(1, sc.team_size + 1)] + ["Summe"]
        mode = f"Summe der besten {sc.team_size}"
    elif sc.team_mode == "durchschnitt":
        heads = ["Summe", "Anzahl", "Durchschnitt"]
        mode = "Durchschnitt aller"
    else:
        heads = ["Summe"]
        mode = f"Summe der besten {sc.team_size}" if sc.team_mode == "beste" else "Summe aller"
    heads.append("Platz")
    label_w = 48 * mm
    cell = min(22 * mm, (total - label_w) / len(heads))
    label = Paragraph(f"<b>{escape(sc.name)}</b><br/>{escape(mode)}", n)
    return Table(
        [
            [label] + [Paragraph(f"<b>{escape(h)}</b>", n) for h in heads],
            [""] + [(filled or {}).get(h, "") for h in heads],
        ],
        colWidths=[label_w] + [cell] * len(heads),
        rowHeights=[None, 9 * mm],
        hAlign="LEFT",
        style=[
            ("SPAN", (0, 0), (0, 1)),
            ("BOX", (0, 0), (-1, -1), 0.8, primary),
            ("GRID", (1, 1), (-1, 1), 0.3, colors.grey),
            ("LINEBELOW", (1, 0), (-1, 0), 0.8, primary),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ],
    )


def shootoff_list(comp: Competition, shootoff) -> bytes:
    """Stechliste: die Betroffenen eines Stechens mit leeren Feldern je Schuss."""
    org = comp.organization
    st = _styles()
    n = st["normal"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    sc = shootoff.scoring
    shots = shootoff.shots if shootoff.shots and shootoff.shots <= 10 else 0
    unit_kind = sc.kind == Scoring.Kind.UNIT
    head = ["Einheit", "Schütze"] if unit_kind else ["Name", "Einheit / Verein"]
    head += [str(i) for i in range(1, shots + 1)] + ["Ergebnis", "Aufsicht"]
    fixed = [50 * mm, 0] + [10 * mm] * shots + [20 * mm, 22 * mm]
    fixed[1] = 172 * mm - sum(fixed)
    rows = [[Paragraph(f"<b>{escape(h)}</b>", n) for h in head]]
    for e in shootoff.entries.select_related("unit", "starter__unit"):
        if unit_kind:
            first, second = e.unit.name if e.unit_id else "", e.name
        else:
            s = e.starter
            first, second = e.name, (s.unit.name if s and s.unit_id else (s.club if s else ""))
        rows.append(
            [Paragraph(escape(first), n), Paragraph(escape(second), n)] + [""] * (len(head) - 2)
        )
    cond = f"{shootoff.shots} Schuss" if shootoff.shots else ""
    if shootoff.counted_shots and shootoff.counted_shots != shootoff.shots:
        cond += f", davon {shootoff.counted_shots} gewertet"
    story = [
        Paragraph(f"<b>{shootoff.number}. Stechen · {escape(sc.name)}</b>", st["title"]),
        Paragraph(escape(" · ".join([comp.title, *_meta(comp)])), n),
        Paragraph(escape(cond), st["small"]),
        Spacer(1, 4 * mm),
        _grid(rows, fixed, primary, height=10 * mm),
    ]
    return _doc(comp, f"Stechliste {comp.title}", story)


def bird_xlsx(comp: Competition) -> bytes:
    from . import bird

    data = bird.state(comp)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pfänder"
    ws.append(["Pfand", "Schütze", "Einheit", "Verein", "Los", "Schuss Nr.", "Schüsse", "Preis"])
    for p in data["prizes"]:
        ws.append(
            [
                safe_cell(p["name"]),
                safe_cell(p["winner_name"]),
                safe_cell(p["unit_name"]),
                safe_cell(p["club"]),
                safe_cell(p["number"]),
                p["shot_number"],
                p["shots_for"],
                safe_cell(p["award"]),
            ]
        )
    _style(ws)
    ws = wb.create_sheet("Schützen")
    ws.append(["Los", "Name", "Einheit", "Verein", "Teilnahme", "gezogen"])
    for s in sorted(
        comp.starters.select_related("unit"),
        key=lambda s: (bird.lot(s.number) is None, bird.lot(s.number) or 0, s.name),
    ):
        ws.append(
            [
                bird.lot(s.number),
                safe_cell(s.name),
                safe_cell(s.unit.name if s.unit_id else ""),
                safe_cell(s.club),
                s.get_presence_display() if s.presence else "",
                {"hut": "Hut", "digital": "digital", "alle": "alle auf einmal"}.get(
                    s.drawn_how, ""
                ),
            ]
        )
    _style(ws)
    ws = wb.create_sheet("Kennzahlen")
    ws.append(["Kennzahl", "Wert"])
    for a, b in _bird_numbers(data):
        ws.append([safe_cell(a), b])
    _style(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx(comp: Competition) -> bytes:
    if comp.form == Competition.Form.BIRD:
        return bird_xlsx(comp)
    data = services.standings(comp)
    scorings = {s.pk: s for s in comp.scorings.all()}
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for block in data["scorings"]:
        sc = scorings[block["scoring"]]
        title = re.sub(r"[\[\]:*?/\\]", "", block["name"])[:28] or "Wertung"
        while title in used:
            title = title[:26] + "_"
        used.add(title)
        ws = wb.create_sheet(title)
        if not comp.own:
            team = sc.kind == Scoring.Kind.TEAM
            ws.append(
                [
                    "Platz",
                    "von",
                    "Mannschaft" if team else "Name",
                    "Schützen" if team else "Einheit",
                ]
                + ["Wert", "Bemerkung"]
            )
            for row in block["rows"]:
                ws.append(
                    [
                        row.get("place"),
                        sc.field_size,
                        safe_cell(row["name"]),
                        safe_cell(", ".join(row.get("members", [])) if team else row["unit_name"]),
                        float(row["value"]) if row.get("value") is not None else None,
                        safe_cell(row.get("note", "")),
                    ]
                )
            _style(ws)
            continue
        if sc.kind == Scoring.Kind.UNIT:
            ws.append(
                ["Platz", "Einheit", "Schützen", "Bester Schütze", "Wert", "Stechen", "Hinweis"]
            )
            for row in block["rows"]:
                ws.append(
                    [
                        row.get("place"),
                        safe_cell(row["name"]),
                        row["count"],
                        safe_cell(row["best"]),
                        float(row["value"]) if row.get("value") is not None else None,
                        " / ".join(row.get("shootoff", {}).values()),
                        safe_cell(row.get("reason", "")),
                    ]
                )
        else:
            ws.append(
                [
                    "Platz",
                    "Name",
                    "Einheit",
                    "Verein",
                    "Gast",
                    "außer Konkurrenz",
                    "Wert",
                    "Stechen",
                ]
            )
            for row in block["rows"]:
                ws.append(
                    [
                        row.get("place"),
                        safe_cell(row["name"]),
                        safe_cell(row["unit_name"]),
                        safe_cell(row["club"]),
                        "ja" if row["guest"] else "",
                        "ja" if row["out_of_competition"] else "",
                        float(row["value"]) if row.get("value") is not None else None,
                        " / ".join(row.get("shootoff", {}).values()),
                    ]
                )
        _style(ws)
    stats = data["stats"]
    if stats:
        ws = wb.create_sheet("Kennzahlen")
        ws.append(["Kennzahl", "Wert"])
        for a, b in (
            ("Schützen", stats["count"]),
            ("Mitglieder", stats["members"]),
            ("Gäste", stats["guests"]),
            ("Durchschnitt", stats["average"]),
            ("Höchster Wert", stats["max"]),
            ("Niedrigster Wert", stats["min"]),
        ):
            ws.append([a, b])
        ws.append([])
        ws.append(
            ["Einheit", "Schützen", "Mitglieder", "Beteiligung %", "Durchschnitt", "Bester Wert"]
        )
        for u in stats["per_unit"]:
            ws.append(
                [
                    safe_cell(u["name"]),
                    u["count"],
                    u["members"],
                    u["participation"],
                    u["average"],
                    u["best"],
                ]
            )
        _style(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style(ws):
    fill = PatternFill("solid", fgColor="DDE7E3")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for i in range(1, ws.max_column + 1):
        width = max((len(str(c.value or "")) for c in ws[get_column_letter(i)]), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)
