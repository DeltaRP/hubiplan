"""Teilnehmerliste eines Termins: Daten für die Ansicht und Export (Excel, PDF)."""

from __future__ import annotations

import io
from dataclasses import dataclass
from xml.sax.saxutils import escape

from django.db.models import Exists
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from exchange.export import safe_cell
from members.models import DepartmentMembership, today
from organizations.models import Department
from organizations.terms import department_terms
from organizations.theme import effective_theme
from organizations.tree import with_descendants

from . import notify, rsvp
from .models import Event, EventResponse


@dataclass
class Row:
    membership_id: int
    name: str
    club_number: str
    departments: str
    response: EventResponse | None

    def data(self) -> dict:
        r = self.response
        return {
            "membership_id": self.membership_id,
            "name": self.name,
            "club_number": self.club_number,
            "departments": self.departments,
            "status": r.status if r else None,
            "guests": r.guests if r else 0,
            "note": r.note if r else "",
            "source": r.source if r else None,
            "updated_at": r.updated_at if r else None,
        }


def _departments(membership_ids: list[int]) -> dict[int, str]:
    on = today()
    names: dict[int, list[str]] = {}
    rows = (
        DepartmentMembership.objects.filter(membership_id__in=membership_ids, valid_from__lte=on)
        .exclude(valid_until__lt=on)
        .select_related("department")
        .order_by("department__name")
    )
    for dm in rows:
        names.setdefault(dm.membership_id, []).append(dm.department.name)
    return {k: ", ".join(v) for k, v in names.items()}


def rows(event: Event, department_id: int | None = None) -> tuple[list[Row], list[Row]]:
    """(Rückmeldungen, noch offen). Mit ``department_id`` nur Mitglieder dieses Zugs
    und aller Einheiten darunter (wie beim Einladen und Eintragen)."""
    members = rsvp.invited(event)
    if department_id is not None:
        org_id = (
            Department.objects.filter(pk=department_id)
            .values_list("organization_id", flat=True)
            .first()
        )
        ids = with_descendants(org_id, [department_id]) if org_id else set()
        ids = list(ids or [department_id])
        members = members.filter(Exists(rsvp._in_department(ids, today())))
    members = list(members)
    responses = {r.membership_id: r for r in event.responses.select_related("membership__person")}
    if department_id is None:
        # Rückmeldungen von inzwischen nicht mehr Eingeladenen (z. B. ausgetreten) bleiben
        known = {m.pk for m in members}
        members += [r.membership for k, r in responses.items() if k not in known]
    depts = _departments([m.pk for m in members])
    answered, open_ = [], []
    for m in members:
        row = Row(
            m.pk,
            f"{m.person.last_name}, {m.person.first_name}",
            m.club_number,
            depts.get(m.pk, ""),
            responses.get(m.pk),
        )
        (answered if row.response else open_).append(row)
    answered.sort(key=lambda r: (r.response.status != EventResponse.Status.YES, r.name))
    return answered, open_


def totals(answered: list[Row]) -> rsvp.Summary:
    """Summen genau der aufgeführten Rückmeldungen (bei einem Zug nur dessen Mitglieder)."""
    responses = [r.response for r in answered if r.response]
    result = rsvp.Summary()
    for r in responses:
        if r.status == EventResponse.Status.YES:
            result.yes += 1
            result.guests += r.guests
        else:
            result.no += 1
    return result


def _dept(event: Event) -> str:
    """Spaltenkopf: eigene Abteilungsbezeichnung des Vereins."""
    return department_terms(event.organization)["sg"]


def _title(event: Event) -> str:
    return f"{event.title} · {notify.when(event)}"


def xlsx(event: Event, answered: list[Row], open_: list[Row]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Teilnehmer"
    ws.append([safe_cell(_title(event))])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    head = ["Name", "Mitgliedsnummer", _dept(event), "Rückmeldung", "Begleitpersonen", "Bemerkung"]
    ws.append([safe_cell(h) for h in head])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    labels = dict(EventResponse.Status.choices)
    for row in answered + open_:
        r = row.response
        # Texte aus Rückmeldungen schreiben Mitglieder selbst: nie als Formel
        ws.append(
            [
                safe_cell(v)
                for v in (
                    row.name,
                    row.club_number,
                    row.departments,
                    labels[r.status] if r else "offen",
                    r.guests if r else None,
                    r.note if r else "",
                )
            ]
        )
    s = totals(answered)
    ws.append([])
    ws.append([f"Zusagen: {s.yes}, Begleitpersonen: {s.guests}, Personen: {s.persons}"])
    ws.append([f"Absagen: {s.no}, offen: {len(open_)}"])
    for i, width in enumerate([30, 16, 22, 14, 16, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pdf(event: Event, theme_chain, answered: list[Row], open_: list[Row]) -> bytes:
    """Liste der Zusagen (zum Abhaken vor Ort), danach Absagen und offene."""
    buf = io.BytesIO()
    theme = effective_theme(theme_chain)
    primary = colors.HexColor(theme["primary"])
    base = getSampleStyleSheet()
    normal = ParagraphStyle("n", parent=base["Normal"], fontSize=9.5, leading=12)
    head_style = ParagraphStyle("h", parent=normal, fontSize=14, leading=18)
    section = ParagraphStyle("m", parent=normal, fontSize=11, leading=14, textColor=primary)
    owner = event.department.name if event.department_id else event.organization.name
    printed = timezone.localtime().strftime("%d.%m.%Y %H:%M")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(18 * mm, 10 * mm, f"{owner} · Teilnehmerliste")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Stand {printed} · Seite {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title=f"Teilnehmerliste {event.title}",
        author=owner,
    )
    s = totals(answered)
    story: list = [
        Paragraph(escape(event.title), head_style),
        Paragraph(escape(f"{owner} · {notify.when(event)}"), normal),
        Paragraph(
            f"Zusagen {s.yes} · Begleitpersonen {s.guests} · Personen {s.persons} · "
            f"Absagen {s.no} · offen {len(open_)}",
            normal,
        ),
        Spacer(1, 5 * mm),
    ]

    def table(title: str, items: list[Row], with_check: bool) -> None:
        if not items:
            return
        data = [["", "Name", _dept(event), "Begl.", "Bemerkung"]]
        for row in items:
            r = row.response
            data.append(
                [
                    "",
                    Paragraph(escape(row.name), normal),
                    Paragraph(escape(row.departments), normal),
                    str(r.guests) if r and r.guests else "",
                    Paragraph(escape(r.note if r else ""), normal),
                ]
            )
        t = Table(data, colWidths=[7 * mm, 55 * mm, 38 * mm, 12 * mm, None], repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8.5),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                ]
            )
        )
        if with_check:
            # Kästchen zum Abhaken vor Ort
            t.setStyle(TableStyle([("GRID", (0, 1), (0, -1), 0.6, colors.black)]))
        story.extend([Paragraph(title, section), Spacer(1, 2 * mm), t, Spacer(1, 5 * mm)])

    yes = [r for r in answered if r.response.status == EventResponse.Status.YES]
    no = [r for r in answered if r.response.status == EventResponse.Status.NO]
    table(f"Zusagen ({len(yes)})", yes, True)
    table(f"Absagen ({len(no)})", no, False)
    table(f"Noch keine Rückmeldung ({len(open_)})", open_, False)
    if not (yes or no or open_):
        story.append(Paragraph("Noch keine Rückmeldungen.", normal))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
