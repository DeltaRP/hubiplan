"""Auswertung einer Veranstaltung bzw. einer Reihe als PDF und Excel."""

from __future__ import annotations

import io
from decimal import Decimal
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from exchange.export import safe_cell
from organizations.theme import effective_theme

from .projects import _comparison, _evaluation
from .report import _euro, _styles, _table


def _period(p: dict) -> str:
    start = p["starts_on"]
    end = p.get("ends_on")
    if end and end != start:
        return f"{start:%d.%m.%Y} bis {end:%d.%m.%Y}"
    return f"{start:%d.%m.%Y}"


def _build(org, title: str, story: list, *, wide: bool = False) -> bytes:
    buf = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(20 * mm, 10 * mm, f"{org.name} · {title}")
        canvas.drawRightString((277 if wide else 190) * mm, 10 * mm, f"Seite {doc.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        buf,
        pagesize=landscape(A4) if wide else A4,
        leftMargin=20 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
        author=org.name,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def project_pdf(project) -> bytes:
    org = project.organization
    data = _evaluation(project)
    st = _styles()
    r, rb, n = st["right"], st["right_bold"], st["normal"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    sub = [escape(org.name), _period(data)]
    if data["department_name"]:
        sub.append(escape(data["department_name"]))
    if data["series_name"]:
        sub.append(f"Reihe {escape(data['series_name'])}")
    story: list = [
        Paragraph(f"<b>Auswertung {escape(data['name'])}</b>", st["title"]),
        Paragraph(" · ".join(sub), n),
        Paragraph(
            f"Stand {timezone.localdate():%d.%m.%Y} · {data['status_label']} · "
            f"{data['booking_count']} Buchungen",
            st["small"],
        ),
        Spacer(1, 4 * mm),
    ]
    t = data["totals"]
    story.append(
        _table(
            [
                [Paragraph("<b>Ergebnis</b>", n), Paragraph("<b>Betrag</b>", r)],
                [Paragraph("Einnahmen", n), Paragraph(_euro(t["income"]), r)],
                [Paragraph("Ausgaben", n), Paragraph(_euro(t["expense"]), r)],
                [
                    Paragraph("<b>Überschuss</b>" if t["result"] >= 0 else "<b>Fehlbetrag</b>", n),
                    Paragraph(_euro(t["result"]), rb),
                ],
            ],
            [80 * mm, 40 * mm],
            primary,
        )
    )
    oi = data["open_invoices"]
    if oi["count"]:
        story.append(
            Paragraph(
                f"Noch offen: {oi['count']} Rechnung(en) über {_euro(oi['outstanding'])} "
                "(nicht im Ergebnis enthalten).",
                st["small"],
            )
        )

    story.append(Paragraph("<b>Nach Kategorie</b>", st["h2"]))
    rows = [
        [
            Paragraph("<b>Kategorie</b>", n),
            Paragraph("<b>Einnahmen</b>", r),
            Paragraph("<b>Ausgaben</b>", r),
            Paragraph("<b>Saldo</b>", r),
        ]
    ]
    for c in data["categories"]:
        rows.append(
            [
                Paragraph(escape(c["name"]), n),
                Paragraph(_euro(c["income"]), r),
                Paragraph(_euro(c["expense"]), r),
                Paragraph(_euro(c["result"]), r),
            ]
        )
    rows.append(
        [
            Paragraph("<b>Summe</b>", n),
            Paragraph(_euro(t["income"]), rb),
            Paragraph(_euro(t["expense"]), rb),
            Paragraph(_euro(t["result"]), rb),
        ]
    )
    story.append(_table(rows, [80 * mm, 30 * mm, 30 * mm, 32 * mm], primary))

    if data["tax_areas_enabled"] and data["tax_areas"]:
        story.append(Paragraph("<b>Nach steuerlichem Bereich</b>", st["h2"]))
        rows = [
            [
                Paragraph("<b>Bereich</b>", n),
                Paragraph("<b>Einnahmen</b>", r),
                Paragraph("<b>Ausgaben</b>", r),
                Paragraph("<b>Saldo</b>", r),
            ]
        ]
        for a in data["tax_areas"]:
            rows.append(
                [
                    Paragraph(escape(a["label"]), n),
                    Paragraph(_euro(a["income"]), r),
                    Paragraph(_euro(a["expense"]), r),
                    Paragraph(_euro(a["result"]), r),
                ]
            )
        story.append(_table(rows, [80 * mm, 30 * mm, 30 * mm, 32 * mm], primary, total_row=False))

    if data["bookings"]:
        story.append(Paragraph("<b>Buchungen</b>", st["h2"]))
        rows = [
            [
                Paragraph(f"<b>{h}</b>", r if h == "Betrag" else n)
                for h in ["Nr.", "Datum", "Buchungstext", "Kategorie", "Betrag"]
            ]
        ]
        for b in data["bookings"]:
            amount = b["amount"] if b["kind"] == "income" else -b["amount"]
            rows.append(
                [
                    Paragraph(escape(b["display_number"]), n),
                    Paragraph(f"{b['date']:%d.%m.%Y}", n),
                    Paragraph(escape(b["text"]), n),
                    Paragraph(escape(b["category_name"] or ""), n),
                    Paragraph(_euro(amount), r),
                ]
            )
        story.append(
            _table(
                rows,
                [24 * mm, 20 * mm, 66 * mm, 36 * mm, 26 * mm],
                primary,
                total_row=False,
            )
        )
    return _build(org, f"Auswertung {data['name']}", story)


def _xlsx(org, sheets: list[tuple[str, list[list]]], money_from: int = 1) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", fgColor="DDE7E3")
    for title, table in sheets:
        ws = wb.create_sheet(title[:31])
        for row in table:
            ws.append([safe_cell(float(v) if isinstance(v, Decimal) else v) for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00 €"
                elif hasattr(cell.value, "year"):
                    cell.number_format = "DD.MM.YYYY"
        for i in range(1, ws.max_column + 1):
            width = max(len(str(r[i - 1])) if i <= len(r) else 0 for r in table[:200]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 50)
        ws.freeze_panes = "A2"
    wb.properties.creator = f"HUBIplan – {org.name}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def project_xlsx(project) -> bytes:
    data = _evaluation(project)
    t = data["totals"]
    summary = [
        ["Kategorie", "Einnahmen", "Ausgaben", "Saldo"],
        *[[c["name"], c["income"], c["expense"], c["result"]] for c in data["categories"]],
        ["Summe", t["income"], t["expense"], t["result"]],
    ]
    bookings = [
        ["Nr.", "Datum", "Art", "Buchungstext", "Kategorie", "Konto", "Beleg-Nr.", "Betrag"],
        *[
            [
                b["display_number"],
                b["date"],
                b["kind_label"],
                b["text"],
                b["category_name"] or "",
                b["account_name"],
                b["reference"],
                b["amount"] if b["kind"] == "income" else -b["amount"],
            ]
            for b in data["bookings"]
        ],
    ]
    sheets = [("Nach Kategorie", summary), ("Buchungen", bookings)]
    if data["tax_areas_enabled"] and data["tax_areas"]:
        sheets.insert(
            1,
            (
                "Steuerliche Bereiche",
                [
                    ["Bereich", "Einnahmen", "Ausgaben", "Saldo"],
                    *[
                        [a["label"], a["income"], a["expense"], a["result"]]
                        for a in data["tax_areas"]
                    ],
                ],
            ),
        )
    return _xlsx(project.organization, sheets)


def _series_table(data: dict) -> list[list]:
    head = ["Kategorie", *[p["name"] for p in data["projects"]]]
    rows = [head]
    for c in data["categories"]:
        rows.append([c["name"], *[v["result"] for v in c["values"]]])
    rows.append(["Einnahmen", *[p["income"] for p in data["projects"]]])
    rows.append(["Ausgaben", *[p["expense"] for p in data["projects"]]])
    rows.append(["Ergebnis", *[p["result"] for p in data["projects"]]])
    return rows


def series_xlsx(series) -> bytes:
    return _xlsx(series.organization, [("Vergleich", _series_table(_comparison(series)))])


def series_pdf(series) -> bytes:
    org = series.organization
    data = _comparison(series)
    st = _styles()
    r, rb, n = st["right"], st["right_bold"], st["normal"]
    primary = colors.HexColor(effective_theme(org.ancestors())["primary"])
    table = _series_table(data)
    count = max(len(data["projects"]), 1)
    wide = count > 3
    usable = (277 if wide else 190) * mm - 20 * mm
    first = 60 * mm
    other = (usable - first) / count
    rows = [
        [Paragraph(f"<b>{escape(str(h))}</b>", n if i == 0 else r) for i, h in enumerate(table[0])]
    ]
    last = len(table) - 1
    for k, row in enumerate(table[1:], start=1):
        bold = k >= last - 2
        rows.append(
            [Paragraph(f"<b>{escape(row[0])}</b>" if bold else escape(row[0]), n)]
            + [Paragraph(_euro(v), rb if k == last else r) for v in row[1:]]
        )
    story = [
        Paragraph(f"<b>Vergleich {escape(data['name'])}</b>", st["title"]),
        Paragraph(
            f"{escape(org.name)} · Saldo je Kategorie · Stand {timezone.localdate():%d.%m.%Y}",
            n,
        ),
        Spacer(1, 4 * mm),
        _table(rows, [first] + [other] * count, primary),
    ]
    return _build(org, f"Vergleich {data['name']}", story, wide=wide)
