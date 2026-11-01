"""Mitglieder- und Jubiläumslisten als Excel oder PDF."""

from __future__ import annotations

import datetime
import io
from collections.abc import Callable
from dataclasses import dataclass
from xml.sax.saxutils import escape

from django.utils import timezone
from django.utils.functional import Promise
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from customfields.models import FieldValue, format_year_range
from customfields.views import definitions_for
from members.duration import jubilees_between
from members.views import filter_members, member_queryset, member_rows
from organizations.models import Organization
from organizations.terms import fill
from organizations.theme import effective_theme

STATUS = {"active": "Mitglied", "former": "ehemalig"}
KIND = {"club": "Verein", "federation": "Dachverband"}


def _date(value) -> str:
    if not value:
        return ""
    if isinstance(value, str):
        value = datetime.date.fromisoformat(value)
    return value.strftime("%d.%m.%Y")


@dataclass
class Column:
    key: str
    label: str
    value: Callable[[dict], object]


@dataclass
class PhotoRef:
    """Platzhalter für ein Foto in einer Exportzeile (nur PDF zeigt das Bild)."""

    file: object | None  # FieldFile des Vorschaubilds (verschlüsselt gespeichert)

    def __str__(self) -> str:
        return "ja" if self.file else ""

    def image_data(self) -> io.BytesIO | None:
        if not self.file:
            return None
        try:
            with self.file.open("rb") as handle:
                return io.BytesIO(handle.read())
        except (OSError, ValueError):
            return None


def member_columns(org: Organization) -> list[Column]:
    cols = [
        Column("photo", "Foto", lambda r: PhotoRef(r.get("photo_path"))),
        Column("last_name", "Name", lambda r: r["last_name"]),
        Column("first_name", "Vorname", lambda r: r["first_name"]),
        Column("birth_date", "Geburtsdatum", lambda r: _date(r["birth_date"])),
        Column("gender", "Geschlecht", lambda r: r.get("gender", "")),
        Column("member_type", "Mitgliedsart", lambda r: r["member_type"] or ""),
        Column("status", "Status", lambda r: STATUS.get(r["status"], r["status"])),
        Column("entry_date", "Eintritt", lambda r: _date(r["entry_date"])),
        Column("exit_date", "Austritt", lambda r: _date(r["exit_date"])),
        Column("club_number", "Mitgliedsnummer", lambda r: r["club_number"]),
        Column("member_number", "Verbandsnummer", lambda r: r["member_number"] or ""),
        Column("departments", org.department_label, lambda r: ", ".join(r["departments"])),
        Column("honorary_titles", "Ehrungen", lambda r: ", ".join(r["honorary_titles"])),
        Column("email", "E-Mail", lambda r: r["email"]),
        Column("phone", "Telefon", lambda r: r["phone"]),
        Column("street", "Straße", lambda r: r["street"]),
        Column("postal_code", "PLZ", lambda r: r["postal_code"]),
        Column("city", "Ort", lambda r: r["city"]),
    ]
    for definition in definitions_for(org):
        if definition.is_archived:
            continue
        cols.append(
            Column(
                f"field:{definition.pk}",
                definition.label,
                (lambda d: lambda r: format_field(d, r["fields"].get(d.pk)))(definition),
            )
        )
    for dignity in dignities_for(org):
        label = f"{dignity.name} ({dignity.unit.name})" if dignity.unit_id else dignity.name
        cols.append(
            Column(
                f"dignity:{dignity.pk}",
                label,
                (lambda pk: lambda r: ", ".join(r.get("reigns", {}).get(pk, [])))(dignity.pk),
            )
        )
    if "competitions" in (org.enabled_modules or []):
        # Titel und Wanderpokale aus den Wettbewerben („Vereinsmeister: 2019, 2023“)
        cols.append(Column("competition_titles", "Wettbewerbstitel", lambda r: r.get("titles", "")))
    return cols


def dignities_for(org: Organization):
    """Würden der Einheit und ihrer Untereinheiten, je eine Exportspalte."""
    from honors.models import Dignity

    return (
        Dignity.objects.filter(organization=org, is_archived=False)
        .select_related("unit")
        .order_by("unit__sort_order", "unit__name", "sort_order", "name")
    )


def reigns_by_person(org: Organization, person_ids) -> dict[int, dict[int, list[str]]]:
    """Je Person und Würde die Jahre der Regentschaften, älteste zuerst."""
    from honors.models import Reign

    reigns = sorted(
        Reign.objects.filter(organization=org, person_id__in=person_ids),
        key=lambda r: (r.sort_year, r.pk),
    )
    result: dict[int, dict[int, list[str]]] = {}
    for reign in reigns:
        result.setdefault(reign.person_id, {}).setdefault(reign.dignity_id, []).append(
            reign.period_text
        )
    return result


def federation_columns(org=None) -> list[Column]:
    club_label = (getattr(org, "department_label", "") or "Verein") if org else "Verein"
    return [
        Column("club", club_label, lambda r: r["club"]),
        Column("last_name", "Name", lambda r: r["last_name"]),
        Column("first_name", "Vorname", lambda r: r["first_name"]),
        Column("birth_date", "Geburtsdatum", lambda r: _date(r["birth_date"])),
        Column("gender", "Geschlecht", lambda r: r.get("gender", "")),
        Column("member_number", "Verbandsnummer", lambda r: r["member_number"] or ""),
        Column("member_type", "Mitgliedsart", lambda r: r["member_type"] or ""),
        Column("honorary_titles", "Ehrungen", lambda r: ", ".join(r["honorary_titles"])),
        Column("status", "Status", lambda r: STATUS.get(r["status"], r["status"])),
        Column("entry_date", "Eintritt", lambda r: _date(r["entry_date"])),
        Column("exit_date", "Austritt", lambda r: _date(r["exit_date"])),
        Column("email", "E-Mail", lambda r: r["email"]),
        Column("phone", "Telefon", lambda r: r["phone"]),
        Column("street", "Straße", lambda r: r["street"]),
        Column("postal_code", "PLZ", lambda r: r["postal_code"]),
        Column("city", "Ort", lambda r: r["city"]),
    ]


def precise_date(value, precision: str) -> str:
    """Datum in der Genauigkeit der Organisation: 15.03.2001, 03/2001 oder 2001."""
    if not value:
        return ""
    if isinstance(value, str):
        value = datetime.date.fromisoformat(value)
    if precision == "year":
        return str(value.year)
    if precision == "month":
        return value.strftime("%m/%Y")
    return value.strftime("%d.%m.%Y")


def jubilee_columns() -> list[Column]:
    return [
        Column("date", "Datum", lambda r: precise_date(r["date"], r.get("precision", "day"))),
        Column("last_name", "Name", lambda r: r["last_name"]),
        Column("first_name", "Vorname", lambda r: r["first_name"]),
        Column("birth_date", "Geburtsdatum", lambda r: _date(r["birth_date"])),
        Column("years", "Jahre", lambda r: r["years"]),
        Column(
            "kind",
            "Jubiläum",
            lambda r: f"{r.get('organization_label') or KIND[r['kind']]}: {r['organization']}",
        ),
        Column("member_number", "Verbandsnummer", lambda r: r["member_number"] or ""),
    ]


def format_field(definition, value) -> str:
    if value in (None, "", []):
        return ""
    if definition.repeatable:
        subs = definition.subfields
        return "; ".join(
            ", ".join(
                _scalar(sub["type"], row.get(sub["key"]))
                for sub in subs
                # 0 und „nein“ sind Werte und bleiben stehen
                if row.get(sub["key"]) not in (None, "")
            )
            for row in value
        )
    return _scalar(definition.field_type, value)


def _scalar(field_type: str, value) -> str:
    if value in (None, ""):
        return ""
    if field_type == "bool":
        return "ja" if value else "nein"
    if field_type == "date":
        return _date(value)
    if field_type == "yearrange":
        return format_year_range(value)
    return str(value)


def member_data(org: Organization, filters: dict) -> list[dict]:
    qs = filter_members(member_queryset(org), filters, org)
    rows = member_rows(list(qs), org, full_names=True)
    person_ids = {
        m.pk: m.person_id for m in qs.only("pk", "person_id")
    }  # gleiche Reihenfolge wie rows
    values: dict[int, dict[int, object]] = {}
    for fv in FieldValue.objects.filter(person_id__in=person_ids.values()):
        values.setdefault(fv.person_id, {})[fv.definition_id] = fv.value
    reigns = reigns_by_person(org, person_ids.values())
    titles = {}
    if "competitions" in (org.enabled_modules or []):
        from competitions.awards import title_summary

        titles = title_summary(org, list(person_ids.keys()))
    for row in rows:
        row["fields"] = values.get(person_ids[row["id"]], {})
        row["reigns"] = reigns.get(person_ids[row["id"]], {})
        row["titles"] = titles.get(row["id"], "")
    from members.photo_models import MemberPhoto

    photos = {}
    for p in MemberPhoto.objects.filter(membership_id__in=person_ids.keys(), status="approved"):
        if p.thumbnail:
            photos[p.membership_id] = p.thumbnail
    for row in rows:
        row["photo_path"] = photos.get(row["id"])
    return rows


def build(job) -> tuple[bytes, str, int]:
    """Erzeugt die Datei eines Exportauftrags: (Inhalt, Dateiname, Zeilen)."""
    org = job.organization
    if job.kind == "federation_members":
        from members.federation_views import (
            federation_queryset,
            federation_rows,
            filter_federation,
        )

        qs = filter_federation(federation_queryset(org), job.filters)
        rows = federation_rows(list(qs), full_names=True)
        all_cols = federation_columns(org)
        title = job.title or fill("Mitglieder {e_im}", org)
    elif job.kind == "jubilees":
        start = datetime.date.fromisoformat(job.filters.get("von"))
        end = datetime.date.fromisoformat(job.filters.get("bis"))
        rows = jubilees_between(org, start, end)
        all_cols = jubilee_columns()
        title = job.title or f"Jubiläen {_date(start)} – {_date(end)}"
    else:
        rows = member_data(org, job.filters)
        all_cols = member_columns(org)
        title = job.title or "Mitgliederliste"
    by_key = {c.key: c for c in all_cols}
    columns = [by_key[k] for k in job.columns if k in by_key] or all_cols[:6]
    table = [[c.label for c in columns]] + [[c.value(r) for c in columns] for r in rows]
    stamp = timezone.localdate().isoformat()
    base = f"{title} {stamp}".replace("/", "-").replace(" ", "_")
    if job.file_format == "pdf":
        content = render_pdf(org, title, table, job.orientation)
        return content, f"{base}.pdf", len(rows)
    return render_xlsx(org, title, table), f"{base}.xlsx", len(rows)


FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Verhindert, dass Excel Text aus den Daten als Formel ausführt."""
    if isinstance(value, (PhotoRef, Promise)):
        return str(value)
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def render_xlsx(org: Organization, title: str, table: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header = table[0]
    ws.append([safe_cell(h) for h in header])
    for row in table[1:]:
        ws.append([safe_cell(v) for v in row])
    fill = PatternFill("solid", fgColor="DDE7E3")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top")
    for i, col in enumerate(header, start=1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in table[1:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.properties.creator = f"HUBIplan – {org.name}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _logo_path(org: Organization) -> str | None:
    for item in org.ancestors():
        if item.logo:
            try:
                return item.logo.path
            except (ValueError, NotImplementedError):
                return None
    return None


def render_pdf(org: Organization, title: str, table: list[list], orientation: str) -> bytes:
    pagesize = landscape(A4) if orientation == "landscape" else A4
    buf = io.BytesIO()
    theme = effective_theme(org.ancestors())
    primary = colors.HexColor(theme["primary"])
    on_primary = colors.HexColor(theme["on_primary"])
    exported = timezone.localtime().strftime("%d.%m.%Y %H:%M")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(15 * mm, 10 * mm, f"{org.name} · exportiert am {exported}")
        canvas.drawRightString(pagesize[0] - 15 * mm, 10 * mm, f"Seite {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=18 * mm,
        title=title,
        author=org.name,
    )
    styles = getSampleStyleSheet()
    story = []
    logo = _logo_path(org)
    heading = [
        Paragraph(f"<b>{escape(org.name)}</b>", styles["Title"]),
        Paragraph(escape(title), styles["Heading2"]),
    ]
    if logo:
        img = Image(logo, width=18 * mm, height=18 * mm, kind="proportional")
        story.append(
            Table(
                [[img, heading]],
                colWidths=[22 * mm, None],
                style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE")],
            )
        )
    else:
        story.extend(heading)
    story.append(Paragraph(f"Stand: {exported} · {len(table) - 1} Einträge", styles["Normal"]))
    story.append(Spacer(1, 5 * mm))
    cell = styles["BodyText"]
    cell.fontSize = 8
    cell.leading = 10
    head = cell.clone("head", textColor=on_primary)
    data = [[Paragraph(f"<b>{escape(str(h))}</b>", head) for h in table[0]]]
    data += [[_pdf_cell(v, cell) for v in row] for row in table[1:]]
    t = Table(data, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), primary),
                ("TEXTCOLOR", (0, 0), (-1, 0), on_primary),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f4f3")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c9cfcc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(t)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _pdf_cell(value, style):
    if isinstance(value, PhotoRef):
        data = value.image_data()
        if data is not None:
            return Image(data, width=12 * mm, height=16 * mm, kind="proportional")
        return ""
    return Paragraph(escape(str(value)) if value is not None else "", style)
