"""Geschlecht als festes Feld der Person (Akte, Liste, Import/Export, Übernahme)."""

import io

import pytest
from django.core.management import call_command
from openpyxl import load_workbook

from circulars.recipients import Candidate, formal_salutation, informal_salutation
from core.tenancy import tenant_context
from customfields.models import FieldDefinition, FieldValue
from members.gender import parse_gender, take_over_field
from members.models import Person

from .test_m2_members import admit, make_types
from .test_m5_exchange import upload

pytestmark = pytest.mark.django_db
JSON = "application/json"


@pytest.fixture(autouse=True)
def _eager(settings, tmp_path):
    settings.CELERY_TASK_ALWAYS_EAGER = True
    settings.PRIVATE_ROOT = tmp_path / "private"


@pytest.fixture
def types(orgs):
    return {key: make_types(orgs[key]) for key in ("a", "b", "c")}


def detail_url(org, pk):
    return f"/api/orgs/{org.pk}/members/{pk}/"


class TestErkennen:
    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("Mann", "männlich"),
            ("männlich", "männlich"),
            (" M ", "männlich"),
            ("Herr", "männlich"),
            ("Frau", "weiblich"),
            ("weiblich", "weiblich"),
            ("w", "weiblich"),
            ("divers", "divers"),
            ("", ""),
            (None, ""),
            ("keine Angabe", ""),
            ("Kind", None),
        ],
    )
    def test_schreibweisen(self, raw, expected):
        assert parse_gender(raw) == expected


class TestAkte:
    def test_aufnahme_und_aendern(self, api, orgs, types, board_a):
        c = api(board_a)
        resp = admit(c, orgs["a"], types["a"]["aktiv"], gender="weiblich")
        assert resp.status_code == 201, resp.json()
        data = resp.json()
        assert data["gender"] == "weiblich"
        assert Person.objects.get().gender == "weiblich"

        resp = c.patch(detail_url(orgs["a"], data["id"]), {"gender": "divers"}, content_type=JSON)
        assert resp.status_code == 200, resp.json()
        assert resp.json()["gender"] == "divers"

        resp = c.patch(detail_url(orgs["a"], data["id"]), {"gender": ""}, content_type=JSON)
        assert resp.status_code == 200
        assert Person.objects.get().gender == ""

    def test_ohne_angabe_und_ungueltig(self, api, orgs, types, board_a):
        c = api(board_a)
        resp = admit(c, orgs["a"], types["a"]["aktiv"])
        assert resp.status_code == 201
        assert resp.json()["gender"] == ""
        resp = c.patch(
            detail_url(orgs["a"], resp.json()["id"]), {"gender": "Mann"}, content_type=JSON
        )
        assert resp.status_code == 400
        assert "gender" in resp.json()

    def test_liste_und_filter(self, api, orgs, types, board_a):
        c = api(board_a)
        admit(c, orgs["a"], types["a"]["aktiv"], last_name="Adler", gender="männlich")
        admit(
            c,
            orgs["a"],
            types["a"]["aktiv"],
            last_name="Berg",
            birth_date="1970-01-01",
            gender="weiblich",
        )
        admit(c, orgs["a"], types["a"]["aktiv"], last_name="Cramer", birth_date="1971-01-01")
        url = f"/api/orgs/{orgs['a'].pk}/members/"

        def names(**params):
            return [r["last_name"] for r in c.get(url, params).json()["results"]]

        assert names() == ["Adler", "Berg", "Cramer"]
        assert names(geschlecht="weiblich") == ["Berg"]
        assert names(geschlecht="männlich") == ["Adler"]
        assert names(geschlecht="ohne") == ["Cramer"]
        rows = c.get(url).json()["results"]
        assert [r["gender"] for r in rows] == ["männlich", "weiblich", ""]

    def test_excel_export(self, api, orgs, types, board_a):
        c = api(board_a)
        admit(c, orgs["a"], types["a"]["aktiv"], last_name="Adler", gender="männlich")
        resp = c.post(
            f"/api/orgs/{orgs['a'].pk}/exports/",
            {"kind": "members", "file_format": "xlsx", "columns": ["last_name", "gender"]},
            content_type=JSON,
        )
        assert resp.status_code == 201, resp.json()
        job = c.get(f"/api/orgs/{orgs['a'].pk}/exports/{resp.json()['id']}/").json()
        file = c.get(f"/api/orgs/{orgs['a'].pk}/exports/{job['id']}/download/")
        ws = load_workbook(io.BytesIO(b"".join(file.streaming_content))).active
        assert [x.value for x in ws[1]] == ["Name", "Geschlecht"]
        assert ws["B2"].value == "männlich"


HEADER = ["Nachname", "Vorname", "Geburtsdatum", "Geschlecht", "Straße", "PLZ", "Ort", "Art"]
HEADER += ["Eintritt"]


def rows(*genders):
    result = [HEADER]
    for i, g in enumerate(genders):
        result.append(
            [f"Person{i}", "Test", f"0{i + 1}.01.1990", g, "Weg 1", "12345", "Stadt", "Akt."]
            + ["01.03.2015"]
        )
    return result


class TestImport:
    def run_import(self, c, org, types, data):
        job = upload(c, org, data).json()
        resp = c.patch(
            f"/api/orgs/{org.pk}/imports/{job['id']}/",
            {
                "mapping": job["mapping"],
                "value_mapping": {"member_type": {"Akt.": str(types["aktiv"].pk)}},
            },
            content_type=JSON,
        )
        assert resp.status_code == 200, resp.json()
        checked = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/").json()
        return job, checked

    def test_mann_frau_werden_erkannt(self, api, orgs, types, board_a):
        c = api(board_a)
        org = orgs["a"]
        # Ein gleichnamiges Zusatzfeld darf das feste Feld nicht verdrängen
        FieldDefinition.objects.create(
            organization=org,
            label="Geschlecht",
            key="geschlecht",
            field_type="choice",
            choices=["männlich", "weiblich"],
        )
        job, checked = self.run_import(c, org, types["a"], rows("Mann", "Frau", ""))
        assert job["mapping"]["3"] == "gender"
        assert checked["counts"]["errors"] == 0, checked
        done = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/").json()
        assert done["status"] == "imported", done
        genders = dict(Person.objects.values_list("last_name", "gender"))
        assert genders == {"Person0": "männlich", "Person1": "weiblich", "Person2": ""}

    def test_unbekannter_wert_ist_fehler(self, api, orgs, types, board_a):
        c = api(board_a)
        _, checked = self.run_import(c, orgs["a"], types["a"], rows("Kind"))
        assert checked["counts"]["errors"] == 1
        assert "Geschlecht" in " ".join(checked["preview"][0]["messages"])


class TestUebernahmeAusZusatzfeld:
    def make_field_values(self, orgs, types, values):
        org = orgs["a"]
        definition = FieldDefinition.objects.create(
            organization=org,
            label="Geschlecht",
            key="geschlecht",
            field_type="choice",
            choices=["männlich", "weiblich", "Mann", "Frau"],
        )
        persons = []
        with tenant_context(org.pk):
            for i, value in enumerate(values):
                p = Person.objects.create(
                    scope=orgs["federation"],
                    last_name=f"P{i}",
                    first_name="Test",
                    birth_date="1990-01-01",
                )
                FieldValue.objects.create(
                    organization=org, definition=definition, person=p, value=value
                )
                persons.append(p)
        return definition, persons

    def test_werte_werden_uebernommen(self, orgs, types):
        definition, persons = self.make_field_values(
            orgs, types, ["Mann", "Frau", "männlich", "weiblich", None]
        )
        # Vorhandene Angabe bleibt
        Person.objects.filter(pk=persons[2].pk).update(gender="divers")
        count = take_over_field(Person, FieldDefinition, FieldValue)
        assert count == 3
        genders = [Person.objects.get(pk=p.pk).gender for p in persons]
        assert genders == ["männlich", "weiblich", "divers", "weiblich", ""]
        # Das Zusatzfeld bleibt unverändert
        assert FieldValue.objects.filter(definition=definition).count() == 5

    def test_befehl(self, orgs, types, capsys):
        self.make_field_values(orgs, types, ["Frau"])
        call_command("geschlecht_uebernehmen")
        assert "1 Personen" in capsys.readouterr().out
        assert Person.objects.get().gender == "weiblich"


class TestAnrede:
    def candidate(self, gender):
        return Candidate(key="m1", club=None, first_name="Max", last_name="Muster", gender=gender)

    def test_anreden(self):
        assert formal_salutation(self.candidate("männlich")) == "Sehr geehrter Herr Muster"
        assert formal_salutation(self.candidate("weiblich")) == "Sehr geehrte Frau Muster"
        assert formal_salutation(self.candidate("")) == "Guten Tag Max Muster"
        assert informal_salutation(self.candidate("männlich")) == "Lieber Max"
        assert informal_salutation(self.candidate("weiblich")) == "Liebe Max"
        assert informal_salutation(self.candidate("divers")) == "Hallo Max"
