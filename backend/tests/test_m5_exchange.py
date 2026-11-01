"""Meilenstein 5: Import von Altdaten und Export von Listen."""

import datetime
import io

import pytest
from openpyxl import Workbook, load_workbook

from access.models import Role, RoleAssignment
from audit.models import AuditEntry
from customfields.models import FieldDefinition, FieldValue
from exchange.models import ImportJob
from members.models import Membership, MemberType, Person

from .factories import audit_with, give_role, make_user
from .test_m2_members import admit, make_types

pytestmark = pytest.mark.django_db
JSON = "application/json"


@pytest.fixture(autouse=True)
def _eager(settings, tmp_path):
    settings.CELERY_TASK_ALWAYS_EAGER = True
    settings.PRIVATE_ROOT = tmp_path / "private"


@pytest.fixture
def types(orgs):
    return {key: make_types(orgs[key]) for key in ("a", "b", "c")}


def xlsx(rows) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "mitglieder.xlsx"
    return buf


HEADER = [
    "Nachname",
    "Vorname",
    "Geburtsdatum",
    "E-Mail",
    "Straße",
    "PLZ",
    "Ort",
    "Art",
    "Eintritt",
    "Austritt",
    "Eintritt 2",
    "Austritt 2",
    "Art 2",
    "Auszeichnung",
    "Datum",
]


def base_rows():
    return [
        HEADER,
        [
            "Muster",
            "Erika",
            "17.05.1980",
            "erika@example.org",
            "Weg 1",
            "12345",
            "Stadt",
            "Akt.",
            "01.03.2015",
            None,
            "01.01.2001",
            "31.12.2005",
            "Pass.",
            "Nadel",
            "2010-05-01",
        ],
        # zweite Zeile derselben Person: weitere Auszeichnung
        [
            "Muster",
            "Erika",
            "17.05.1980",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "Pokal",
            "2018-07-14",
        ],
        ["Fehler", "Fritz", "kein Datum", "f@example.org", "Weg", "1", "x", "Akt.", "2020"],
        [
            "Alt",
            "Anton",
            datetime.date(1950, 1, 1),
            None,
            None,
            None,
            None,
            "Pass.",
            "1970",
            "1999",
        ],
    ]


def upload(client, org, rows):
    return client.post(f"/api/orgs/{org.pk}/imports/", {"file": xlsx(rows)})


@pytest.fixture
def awards(orgs):
    return FieldDefinition.objects.create(
        organization=orgs["a"],
        label="Auszeichnung",
        key="auszeichnung",
        repeatable=True,
        subfields=[
            {"key": "bezeichnung", "label": "Bezeichnung", "type": "text", "choices": []},
            {"key": "datum", "label": "Datum", "type": "date", "choices": []},
        ],
    )


def configure(c, org, job, types, awards):
    mapping = dict(job["mapping"])
    mapping["13"] = f"field:{awards.pk}:bezeichnung"
    mapping["14"] = f"field:{awards.pk}:datum"
    resp = c.patch(
        f"/api/orgs/{org.pk}/imports/{job['id']}/",
        {
            "mapping": mapping,
            "value_mapping": {
                "member_type": {"Akt.": str(types["aktiv"].pk), "Pass.": str(types["passiv"].pk)}
            },
        },
        content_type=JSON,
    )
    assert resp.status_code == 200, resp.json()
    return resp.json()


class TestImport:
    def test_vorschlag_der_zuordnung(self, api, orgs, types, board_a):
        resp = upload(api(board_a), orgs["a"], base_rows())
        assert resp.status_code == 201, resp.json()
        job = resp.json()
        assert job["row_count"] == 4
        m = job["mapping"]
        assert (m["0"], m["1"], m["2"], m["7"], m["8"], m["10"], m["12"]) == (
            "last_name",
            "first_name",
            "birth_date",
            "member_type",
            "entry_date",
            "entry_date_2",
            "member_type_2",
        )

    def test_werte_zuordnen_pruefen_importieren(self, api, orgs, types, board_a, awards):
        c = api(board_a)
        org = orgs["a"]
        job = upload(c, org, base_rows()).json()
        values = c.get(f"/api/orgs/{org.pk}/imports/{job['id']}/values/", {"ziel": "member_type"})
        assert values.json() == ["Akt.", "Pass."]
        configure(c, org, job, types["a"], awards)

        checked = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/").json()
        assert checked["status"] == "validated"
        assert checked["counts"] == {"persons": 3, "ok": 2, "errors": 1, "duplicates": 0, "rows": 4}
        bad = next(p for p in checked["preview"] if p["status"] == "error")
        assert bad["rows"] == [4] and "Geburtsdatum" in bad["messages"][0]
        assert not Membership.objects.exists()  # Prüfen speichert nichts

        done = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/").json()
        assert done["status"] == "imported", done
        assert done["counts"]["imported"] == 2
        assert done["has_error_file"] is True
        erika = Membership.objects.get(person__last_name="Muster")
        periods = list(erika.periods.values_list("entry_date", "exit_date"))
        assert periods == [
            (datetime.date(2001, 1, 1), datetime.date(2005, 12, 31)),
            (datetime.date(2015, 3, 1), None),
        ]
        assert erika.current_type().member_type.name == "aktiv"
        assert erika.person.federation.member_number == 1
        value = FieldValue.objects.get(person=erika.person).value
        assert [v["bezeichnung"] for v in value] == ["Nadel", "Pokal"]
        anton = Membership.objects.get(person__last_name="Alt")
        assert anton.current_period() is None  # ehemaliges Mitglied
        # Fehlerliste herunterladen
        errors = c.get(f"/api/orgs/{org.pk}/imports/{job['id']}/errors/")
        wb = load_workbook(io.BytesIO(b"".join(errors.streaming_content)))
        assert wb.active["B2"].value == "Fehler"
        assert len(audit_with("Import")) == 2

    def test_rueckgaengig(self, api, orgs, types, board_a, awards):
        c = api(board_a)
        org = orgs["a"]
        job = upload(c, org, base_rows()).json()
        configure(c, org, job, types["a"], awards)
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/")
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/")
        resp = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/undo/")
        assert resp.status_code == 200, resp.json()
        assert resp.json()["status"] == "undone"
        assert not Membership.objects.exists()
        assert not Person.objects.exists()
        assert not FieldValue.objects.exists()

    def test_rueckgaengig_blockiert_nach_bearbeitung(self, api, orgs, types, board_a, awards):
        c = api(board_a)
        org = orgs["a"]
        job = upload(c, org, base_rows()).json()
        configure(c, org, job, types["a"], awards)
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/")
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/")
        erika = Membership.objects.get(person__last_name="Muster")
        c.patch(f"/api/orgs/{org.pk}/members/{erika.pk}/", {"city": "Neu"}, content_type=JSON)
        job_data = c.get(f"/api/orgs/{org.pk}/imports/{job['id']}/").json()
        assert job_data["undo_blockers"]
        assert c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/undo/").status_code == 400
        assert Membership.objects.count() == 2

    def test_dublette_im_dachverband_verknuepft(self, api, orgs, types, board_a, board_b, awards):
        admit(api(board_b), orgs["b"], types["b"]["passiv"])  # Erika ist schon in Verein B
        c = api(board_a)
        org = orgs["a"]
        job = upload(c, org, base_rows()[:2]).json()
        configure(c, org, job, types["a"], awards)
        checked = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/").json()
        assert checked["preview"][0]["status"] == "duplicate"
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/")
        assert Person.objects.count() == 1
        assert Membership.objects.count() == 2

    def test_csv_mit_semikolon(self, api, orgs, types, board_a):
        content = (
            "Name;Vorname;Geburtsdatum;Mitgliedsart;Eintritt\nKurz;Karl;1990-02-02;aktiv;2010\n"
        )
        f = io.BytesIO(content.encode("cp1252"))
        f.name = "alt.csv"
        resp = api(board_a).post(f"/api/orgs/{orgs['a'].pk}/imports/", {"file": f})
        assert resp.status_code == 201, resp.json()
        job = resp.json()
        assert job["headers"] == ["Name", "Vorname", "Geburtsdatum", "Mitgliedsart", "Eintritt"]
        checked = (
            api(board_a).post(f"/api/orgs/{orgs['a'].pk}/imports/{job['id']}/validate/").json()
        )
        # Ehemalige ohne Kontakt sind ok; aktuelle brauchen eine Anschrift (E-Mail freiwillig)
        assert checked["counts"]["errors"] == 1
        assert "Anschrift Pflicht" in checked["preview"][0]["messages"][0]

    def test_vorstandsrollen_historie(self, api, orgs, types, board_a):
        role = Role.objects.get(organization=orgs["a"], name="Vorstand")
        rows = [
            [
                "Name",
                "Vorname",
                "Geburtsdatum",
                "Mitgliedsart",
                "Eintritt",
                "Austritt",
                "Vorstandsrolle",
                "Vorstandsrolle von",
                "Vorstandsrolle bis",
            ],
            ["Früher", "Franz", "1940-01-01", "passiv", "1960", "2000", "Vorstand", "1980", "1990"],
        ]
        c = api(board_a)
        job = upload(c, orgs["a"], rows).json()
        mapping = {
            "0": "last_name",
            "1": "first_name",
            "2": "birth_date",
            "3": "member_type",
            "4": "entry_date",
            "5": "exit_date",
            "6": "board_role",
            "7": "board_role_from",
            "8": "board_role_until",
        }
        c.patch(
            f"/api/orgs/{orgs['a'].pk}/imports/{job['id']}/",
            {"mapping": mapping},
            content_type=JSON,
        )
        c.post(f"/api/orgs/{orgs['a'].pk}/imports/{job['id']}/validate/")
        c.post(f"/api/orgs/{orgs['a'].pk}/imports/{job['id']}/run/")
        ra = RoleAssignment.objects.get(role=role, user__isnull=True)
        assert ra.person.last_name == "Früher"
        assert ra.valid_until == datetime.date(1990, 1, 1)
        # Zählt nicht als aktive Rolle eines Logins
        assert api(board_a).get(f"/api/orgs/{orgs['a'].pk}/assignments/").status_code == 200

    def test_nur_mit_bearbeitungsrecht(self, api, orgs, types):
        user = make_user("leser@example.org")
        give_role(user, orgs["a"], ["members.view"], name="Leser")
        assert upload(api(user), orgs["a"], base_rows()).status_code == 403

    def test_fremder_import_unsichtbar(self, api, orgs, types, board_a, board_b):
        job = upload(api(board_a), orgs["a"], base_rows()).json()
        assert api(board_b).get(f"/api/orgs/{orgs['b'].pk}/imports/{job['id']}/").status_code == 404
        assert ImportJob.objects.count() == 1


class TestExport:
    @pytest.fixture
    def members(self, api, orgs, types, board_a):
        c = api(board_a)
        admit(c, orgs["a"], types["a"]["aktiv"], last_name="Adler")
        admit(c, orgs["a"], types["a"]["passiv"], last_name="Berg", birth_date="1970-01-01")
        return c

    def test_excel_mit_spaltenauswahl_und_filter(self, orgs, types, members):
        c = members
        resp = c.post(
            f"/api/orgs/{orgs['a'].pk}/exports/",
            {
                "kind": "members",
                "file_format": "xlsx",
                "columns": ["last_name", "first_name", "member_type", "email"],
                "filters": {"art": types["a"]["aktiv"].pk, "status": "aktiv"},
            },
            content_type=JSON,
        )
        assert resp.status_code == 201, resp.json()
        job = c.get(f"/api/orgs/{orgs['a'].pk}/exports/{resp.json()['id']}/").json()
        assert job["status"] == "done", job
        assert job["row_count"] == 1
        file = c.get(f"/api/orgs/{orgs['a'].pk}/exports/{job['id']}/download/")
        ws = load_workbook(io.BytesIO(b"".join(file.streaming_content))).active
        assert [c.value for c in ws[1]] == ["Name", "Vorname", "Mitgliedsart", "E-Mail"]
        assert ws["A2"].value == "Adler"
        entry = AuditEntry.objects.filter(object_type="Export").get()
        assert "Mitgliedsart" in entry.changes["Spalten"][1]

    def test_pdf_quer(self, orgs, members):
        resp = members.post(
            f"/api/orgs/{orgs['a'].pk}/exports/",
            {
                "kind": "members",
                "file_format": "pdf",
                "orientation": "landscape",
                "columns": ["last_name", "first_name", "birth_date"],
                "filters": {"status": "alle"},
            },
            content_type=JSON,
        )
        job = members.get(f"/api/orgs/{orgs['a'].pk}/exports/{resp.json()['id']}/").json()
        assert job["status"] == "done", job
        pdf = b"".join(
            members.get(f"/api/orgs/{orgs['a'].pk}/exports/{job['id']}/download/").streaming_content
        )
        assert pdf.startswith(b"%PDF")

    def test_jubilaeumsliste(self, orgs, members):
        resp = members.post(
            f"/api/orgs/{orgs['a'].pk}/exports/",
            {
                "kind": "jubilees",
                "file_format": "xlsx",
                "columns": ["date", "last_name", "years"],
                "filters": {"von": "2025-01-01", "bis": "2025-12-31"},
            },
            content_type=JSON,
        )
        assert resp.status_code == 201, resp.json()
        job = members.get(f"/api/orgs/{orgs['a'].pk}/exports/{resp.json()['id']}/").json()
        assert job["status"] == "done"
        # 10 Jahre (Eintritt 2015) für beide – im Verein und im Dachverband
        assert job["row_count"] == 4

    def test_export_braucht_recht(self, api, orgs, types):
        user = make_user("leser@example.org")
        give_role(user, orgs["a"], ["members.view"], name="Leser")
        resp = api(user).post(
            f"/api/orgs/{orgs['a'].pk}/exports/",
            {"kind": "members", "file_format": "xlsx"},
            content_type=JSON,
        )
        assert resp.status_code == 403

    def test_unbekannte_spalte(self, orgs, members):
        resp = members.post(
            f"/api/orgs/{orgs['a'].pk}/exports/",
            {
                "kind": "members",
                "file_format": "xlsx",
                "columns": ["passwort"],
            },
            content_type=JSON,
        )
        assert resp.status_code == 400

    def test_vorlagen(self, orgs, members):
        url = f"/api/orgs/{orgs['a'].pk}/export-templates/"
        resp = members.post(
            url,
            {"name": "Geburtstagsliste", "columns": ["last_name", "birth_date"]},
            content_type=JSON,
        )
        assert resp.status_code == 201
        assert members.get(url).json()[0]["name"] == "Geburtstagsliste"

    def test_spalten_mit_zusatzfeldern(self, orgs, members):
        FieldDefinition.objects.create(
            organization=orgs["a"], label="Schuhgröße", key="schuh", field_type="number"
        )
        cols = members.get(f"/api/orgs/{orgs['a'].pk}/export-columns/").json()["members"]
        assert any(c["label"] == "Schuhgröße" for c in cols)
        assert MemberType.objects.count() == 9


class TestImportWidersprueche:
    """Fachfrage 11 (Audit 31.10.): widersprüchliche Zeilen derselben Person nicht still
    übergehen, sondern in der Vorschau nennen."""

    def test_abweichende_mail_ist_fehler(self, api, orgs, types, board_a, awards):
        c = api(board_a)
        org = orgs["a"]
        rows = base_rows()
        rows[2][3] = "andere@example.org"  # zweite Zeile von Erika mit anderer Mail
        job = upload(c, org, rows).json()
        configure(c, org, job, types["a"], awards)
        checked = c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/").json()
        erika = next(p for p in checked["preview"] if p["rows"] == [2, 3])
        assert erika["status"] == "error"
        assert any("E-Mail" in m and "Zeile 2" in m for m in erika["messages"])
