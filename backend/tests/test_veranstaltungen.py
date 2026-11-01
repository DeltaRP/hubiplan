"""Veranstaltungen: Buchungen und Rechnungen zuordnen, Termine verknüpfen, Auswertung."""

import datetime
import io
import uuid

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from core.tenancy import tenant_context
from events.models import Event
from finance.models import Booking, Project

from .factories import audit_with, give_role, make_member, make_user

pytestmark = pytest.mark.django_db
JSON = "application/json"


@pytest.fixture(autouse=True)
def _private(settings, tmp_path):
    settings.PRIVATE_ROOT = tmp_path / "private"


@pytest.fixture
def club(orgs):
    club = orgs["a"]
    club.enabled_modules = ["finance", "members", "billing", "events"]
    club.save()
    return club


@pytest.fixture
def fin(club):
    return f"/api/orgs/{club.pk}/finance"


@pytest.fixture
def c(api, board_a):
    return api(board_a)


@pytest.fixture
def kasse(c, fin):
    bank = c.post(
        f"{fin}/accounts/",
        {"name": "Girokonto", "opening_date": "2024-01-01"},
        content_type=JSON,
    ).json()["id"]
    cash = c.post(
        f"{fin}/accounts/",
        {"name": "Barkasse", "kind": "cash", "opening_date": "2024-01-01"},
        content_type=JSON,
    ).json()["id"]
    cats = {
        name: c.post(f"{fin}/categories/", {"name": name}, content_type=JSON).json()["id"]
        for name in ("Getränke", "Musik", "Eintritt")
    }
    return {"bank": bank, "cash": cash, **cats}


@pytest.fixture
def fest(c, fin):
    series = c.post(f"{fin}/series/", {"name": "Schützenfest"}, content_type=JSON)
    assert series.status_code == 201, series.json()
    resp = c.post(
        f"{fin}/projects/",
        {
            "name": "Schützenfest 2025",
            "starts_on": "2025-08-23",
            "ends_on": "2025-08-25",
            "series": series.json()["id"],
        },
        content_type=JSON,
    )
    assert resp.status_code == 201, resp.json()
    return resp.json()


def book(c, fin, kasse, project=None, **data):
    payload = {
        "date": "2025-08-23",
        "kind": "income",
        "amount": "100.00",
        "account": kasse["bank"],
        "category": kasse["Eintritt"],
        "text": "Buchung",
        "project": project,
        **data,
    }
    resp = c.post(f"{fin}/bookings/", payload, content_type=JSON)
    assert resp.status_code == 201, resp.json()
    return resp.json()


class TestVeranstaltung:
    def test_anlegen_buchen_und_auswerten(self, c, fin, kasse, fest):
        pid = fest["id"]
        book(c, fin, kasse, pid, amount="800.00", category=kasse["Eintritt"], text="Eintritt")
        book(c, fin, kasse, pid, amount="1200.50", category=kasse["Getränke"], text="Theke")
        book(
            c,
            fin,
            kasse,
            pid,
            kind="expense",
            amount="900",
            category=kasse["Getränke"],
            text="Brauerei",
        )
        book(c, fin, kasse, pid, kind="expense", amount="1500", category=kasse["Musik"])
        # nicht zur Veranstaltung: zählt nicht
        book(c, fin, kasse, None, amount="50")
        umbuchung = book(
            c, fin, kasse, pid, kind="transfer", to_account=kasse["cash"], category=None
        )
        assert umbuchung["project"] is None

        listing = c.get(f"{fin}/projects/").json()
        assert listing[0]["income"] == "2000.50"
        assert listing[0]["expense"] == "2400.00"
        assert listing[0]["result"] == "-399.50"
        assert listing[0]["booking_count"] == 4
        assert listing[0]["series_name"] == "Schützenfest"

        ev = c.get(f"{fin}/projects/{pid}/evaluation/").json()
        cats = {row["name"]: row for row in ev["categories"]}
        assert cats["Getränke"]["income"] == "1200.50"
        assert cats["Getränke"]["expense"] == "900.00"
        assert cats["Getränke"]["result"] == "300.50"
        assert cats["Musik"]["result"] == "-1500.00"
        assert ev["totals"]["result"] == "-399.50"
        assert len(ev["bookings"]) == 4

        # Im Kassenbuch ganz normal, filterbar nach Veranstaltung
        all_bookings = c.get(f"{fin}/bookings/?jahr=2025").json()
        assert all_bookings["count"] == 6
        only = c.get(f"{fin}/bookings/?veranstaltung={pid}").json()
        assert only["count"] == 4
        assert only["results"][0]["project_name"] == "Schützenfest 2025"
        without = c.get(f"{fin}/bookings/?veranstaltung=-").json()
        assert without["count"] == 2

    def test_ende_vor_beginn(self, c, fin):
        resp = c.post(
            f"{fin}/projects/",
            {"name": "X", "starts_on": "2025-05-02", "ends_on": "2025-05-01"},
            content_type=JSON,
        )
        assert resp.status_code == 400
        assert "ends_on" in resp.json()

    def test_zuordnen_auch_im_abgeschlossenen_jahr(self, c, fin, kasse, club):
        old = book(c, fin, kasse, None, date="2024-06-01", text="Altes Fest")
        project = c.post(
            f"{fin}/projects/", {"name": "Fest 2024", "starts_on": "2024-06-01"}, content_type=JSON
        ).json()
        assert c.post(f"{fin}/closings/", {"year": 2024}, content_type=JSON).status_code == 201

        # Betrag bleibt gesperrt, die Veranstaltung lässt sich setzen
        resp = c.patch(f"{fin}/bookings/{old['id']}/", {"amount": "5"}, content_type=JSON)
        assert resp.status_code == 400
        resp = c.patch(
            f"{fin}/bookings/{old['id']}/", {"project": project["id"]}, content_type=JSON
        )
        assert resp.status_code == 200, resp.json()
        assert resp.json()["project"] == project["id"]

        resp = c.post(
            f"{fin}/projects/{project['id']}/unassign/",
            {"bookings": [old["id"]]},
            content_type=JSON,
        )
        assert resp.json() == {"changed": 1}
        resp = c.post(
            f"{fin}/projects/{project['id']}/assign/",
            {"bookings": [old["id"]]},
            content_type=JSON,
        )
        assert resp.json() == {"changed": 1}
        assert Booking.objects.get(pk=old["id"]).project_id == project["id"]
        assert audit_with("Veranstaltung")

    def test_zeitraum_filter_im_kassenbuch(self, c, fin, kasse):
        book(c, fin, kasse, None, date="2025-08-01")
        book(c, fin, kasse, None, date="2025-08-24")
        resp = c.get(f"{fin}/bookings/?von=2025-08-20&bis=2025-08-31").json()
        assert resp["count"] == 1

    def test_loeschen_nur_ohne_buchungen(self, c, fin, kasse, fest, club):
        b = book(c, fin, kasse, fest["id"])
        resp = c.delete(f"{fin}/projects/{fest['id']}/")
        assert resp.status_code == 409
        assert "Buchung" in resp.json()["detail"]
        c.post(f"{fin}/projects/{fest['id']}/unassign/", {"bookings": [b["id"]]}, content_type=JSON)
        with tenant_context(club.pk):
            event = Event.objects.create(
                organization=club,
                title="Kirmes",
                starts_at=timezone.now(),
                ends_at=timezone.now(),
                project_id=fest["id"],
            )
        assert c.delete(f"{fin}/projects/{fest['id']}/").status_code == 204
        event.refresh_from_db()
        assert event.project_id is None  # Termin bleibt im Kalender

    def test_reihe_vergleich(self, c, fin, kasse, fest):
        book(c, fin, kasse, fest["id"], amount="300", category=kasse["Getränke"])
        fest26 = c.post(
            f"{fin}/projects/",
            {"name": "Schützenfest 2026", "starts_on": "2026-08-22", "series": fest["series"]},
            content_type=JSON,
        ).json()
        book(
            c, fin, kasse, fest26["id"], date="2026-08-22", amount="450", category=kasse["Getränke"]
        )
        book(
            c,
            fin,
            kasse,
            fest26["id"],
            date="2026-08-22",
            kind="expense",
            amount="100",
            category=kasse["Musik"],
        )
        data = c.get(f"{fin}/series/{fest['series']}/comparison/").json()
        assert [p["name"] for p in data["projects"]] == ["Schützenfest 2025", "Schützenfest 2026"]
        assert [p["result"] for p in data["projects"]] == ["300.00", "350.00"]
        drinks = next(row for row in data["categories"] if row["name"] == "Getränke")
        assert [v["result"] for v in drinks["values"]] == ["300.00", "450.00"]
        music = next(row for row in data["categories"] if row["name"] == "Musik")
        assert [v["result"] for v in music["values"]] == ["0.00", "-100.00"]
        assert c.get(f"{fin}/projects/?reihe={fest['series']}").json()[0]["name"] == (
            "Schützenfest 2026"
        )
        # Reihe löschen lässt die Veranstaltungen stehen
        assert c.delete(f"{fin}/series/{fest['series']}/").status_code == 204
        assert Project.objects.filter(pk=fest26["id"], series__isnull=True).exists()

    def test_berichte(self, c, fin, kasse, fest):
        book(c, fin, kasse, fest["id"], text="=SUMME(A1)")
        pdf = c.get(f"{fin}/projects/{fest['id']}/report/")
        assert pdf.status_code == 200
        assert pdf.content.startswith(b"%PDF")
        xlsx = c.get(f"{fin}/projects/{fest['id']}/report/?als=xlsx")
        assert xlsx.status_code == 200
        wb = load_workbook(io.BytesIO(xlsx.content))
        assert wb.sheetnames == ["Nach Kategorie", "Buchungen"]
        rows = list(wb["Buchungen"].values)
        assert rows[1][3] == "'=SUMME(A1)"  # keine Formel aus Daten
        assert rows[1][7] == 100.0
        series_pdf = c.get(f"{fin}/series/{fest['series']}/report/")
        assert series_pdf.content.startswith(b"%PDF")
        series_xlsx = c.get(f"{fin}/series/{fest['series']}/report/?als=xlsx")
        assert load_workbook(io.BytesIO(series_xlsx.content)).sheetnames == ["Vergleich"]


class TestTermine:
    def test_termine_verknuepfen(self, c, fin, fest, club):
        start = timezone.make_aware(datetime.datetime(2025, 8, 23, 18))
        with tenant_context(club.pk):
            samstag = Event.objects.create(
                organization=club, title="Kirmes Samstag", starts_at=start, ends_at=start
            )
            sonntag = Event.objects.create(
                organization=club,
                title="Kirmes Sonntag",
                starts_at=start + datetime.timedelta(days=1),
                ends_at=start + datetime.timedelta(days=1),
            )
            weit_weg = Event.objects.create(
                organization=club,
                title="Weihnachtsfeier",
                starts_at=start + datetime.timedelta(days=120),
                ends_at=start + datetime.timedelta(days=120),
            )
        candidates = c.get(f"{fin}/projects/{fest['id']}/event-candidates/").json()
        assert {e["id"] for e in candidates} == {samstag.pk, sonntag.pk}
        found = c.get(f"{fin}/projects/{fest['id']}/event-candidates/?suche=weih").json()
        assert [e["id"] for e in found] == [weit_weg.pk]

        resp = c.post(
            f"{fin}/projects/{fest['id']}/link-events/",
            {"events": [samstag.pk, sonntag.pk]},
            content_type=JSON,
        )
        assert resp.status_code == 200, resp.json()
        assert [e["title"] for e in resp.json()["events"]] == ["Kirmes Samstag", "Kirmes Sonntag"]
        assert c.get(f"{fin}/projects/").json()[0]["event_count"] == 2

        other = c.post(
            f"{fin}/projects/", {"name": "Anderes", "starts_on": "2025-08-23"}, content_type=JSON
        ).json()
        resp = c.post(
            f"{fin}/projects/{other['id']}/link-events/",
            {"events": [samstag.pk]},
            content_type=JSON,
        )
        assert resp.status_code == 400
        resp = c.post(
            f"{fin}/projects/{fest['id']}/link-events/",
            {"unlink": [sonntag.pk]},
            content_type=JSON,
        )
        assert [e["title"] for e in resp.json()["events"]] == ["Kirmes Samstag"]


class TestRechnungen:
    @pytest.fixture
    def billing(self, c, club, kasse):
        base = f"/api/orgs/{club.pk}/billing"
        resp = c.patch(
            f"{base}/settings/",
            {"account": kasse["bank"], "category": kasse["Eintritt"]},
            content_type=JSON,
        )
        assert resp.status_code == 200, resp.json()
        return base

    def test_zahlung_uebernimmt_veranstaltung(self, c, club, fin, billing, fest):
        member = make_member(club, "Erika", "Muster")
        inv = c.post(
            f"{billing}/invoices/",
            {
                "recipient_kind": "member",
                "membership": member.pk,
                "subject": "Standmiete Schützenfest",
                "project": fest["id"],
                "lines": [{"text": "Standmiete", "amount": "80.00"}],
            },
            content_type=JSON,
        ).json()
        assert inv["project"] == fest["id"]
        assert inv["project_name"] == "Schützenfest 2025"
        assert c.post(f"{billing}/invoices/{inv['id']}/finalize/").status_code == 200
        resp = c.post(
            f"{billing}/invoices/{inv['id']}/payments/",
            {
                "date": "2025-08-20",
                "amount": "80.00",
                "method": "transfer",
                "account": c.get(f"{fin}/accounts/").json()[0]["id"],
            },
            content_type=JSON,
            HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()),
        )
        assert resp.status_code == 200, resp.json()
        booking = Booking.objects.get(payment__invoice_id=inv["id"])
        assert booking.project_id == fest["id"]
        assert c.get(f"{billing}/invoices/?veranstaltung={fest['id']}").json()["count"] == 1

        # Festgeschrieben: Veranstaltung noch änderbar, die Buchung zieht mit
        other = c.post(
            f"{fin}/projects/", {"name": "Herbstfest", "starts_on": "2025-10-01"}, content_type=JSON
        ).json()
        resp = c.patch(
            f"{billing}/invoices/{inv['id']}/", {"project": other["id"]}, content_type=JSON
        )
        assert resp.status_code == 200, resp.json()
        booking.refresh_from_db()
        assert booking.project_id == other["id"]
        # Rechnung hängt an der Veranstaltung: nicht löschbar
        assert c.delete(f"{fin}/projects/{other['id']}/").status_code == 409

    def test_offene_rechnungen_in_auswertung(self, c, club, billing, fin, fest):
        member = make_member(club, "Erika", "Muster")
        inv = c.post(
            f"{billing}/invoices/",
            {
                "recipient_kind": "member",
                "membership": member.pk,
                "project": fest["id"],
                "lines": [{"text": "Standmiete", "amount": "80.00"}],
            },
            content_type=JSON,
        ).json()
        c.post(f"{billing}/invoices/{inv['id']}/finalize/")
        ev = c.get(f"{fin}/projects/{fest['id']}/evaluation/").json()
        assert ev["open_invoices"] == {"count": 1, "outstanding": "80.00"}


class TestRechte:
    def test_nur_mit_finanzrechten(self, c, fin, club, api, fest):
        pruefer = make_user("pruefer@example.org")
        give_role(pruefer, club, ["finance.view"], name="Kassenprüfer")
        p = api(pruefer)
        assert p.get(f"{fin}/projects/").status_code == 200
        assert p.get(f"{fin}/projects/{fest['id']}/evaluation/").status_code == 200
        assert (
            p.post(f"{fin}/projects/", {"name": "X", "starts_on": "2025-01-01"}).status_code == 403
        )
        assert p.post(f"{fin}/projects/{fest['id']}/assign/", {"bookings": []}).status_code == 403
        jemand = make_user("schrift@example.org")
        give_role(jemand, club, ["members.view"], name="Schriftführer")
        assert api(jemand).get(f"{fin}/projects/").status_code == 403

    def test_fremde_veranstaltung(self, c, fin, kasse, fest, orgs, api, board_b):
        club_b = orgs["b"]
        club_b.enabled_modules = ["finance"]
        club_b.save()
        b = api(board_b)
        base_b = f"/api/orgs/{club_b.pk}/finance"
        assert b.get(f"{base_b}/projects/").json() == []
        assert b.get(f"{base_b}/projects/{fest['id']}/evaluation/").status_code == 404
        # Buchung mit fremder Veranstaltung
        resp = c.post(
            f"{fin}/projects/", {"name": "Eigenes", "starts_on": "2025-01-01"}, content_type=JSON
        )
        assert resp.status_code == 201
        acc = b.post(
            f"{base_b}/accounts/",
            {"name": "Konto", "opening_date": "2024-01-01"},
            content_type=JSON,
        ).json()["id"]
        cat = b.post(f"{base_b}/categories/", {"name": "Spenden"}, content_type=JSON).json()["id"]
        resp = b.post(
            f"{base_b}/bookings/",
            {
                "date": "2025-02-01",
                "kind": "income",
                "amount": "5",
                "account": acc,
                "category": cat,
                "text": "x",
                "project": fest["id"],
            },
            content_type=JSON,
        )
        assert resp.status_code == 400
        assert "project" in resp.json()
