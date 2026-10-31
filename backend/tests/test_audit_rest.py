# ruff: noqa: F811  (Fixtures aus anderen Testmodulen)
"""Regressionstests zum Datenaudit vom 25.10.2026: Termine, Statistik, Zusatzfelder,
Import und Nachverarbeitung (MEM-E01–E08, DS04/05/07/09/10, K03)."""

import datetime
import io
from unittest import mock

import pytest
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook

from customfields.models import FieldDefinition, clean_scalar
from events import recurrence
from events.models import Event, EventException, EventResponse
from exchange import tasks
from exchange.export import format_field
from exchange.models import ImportJob
from members.duration import jubilee_date
from members.models import ContactDetails, MembershipPeriod

from .factories import make_member
from .test_m5_exchange import _eager, awards, base_rows, configure, types, upload  # noqa: F401
from .test_m12_events import (  # noqa: F401
    JSON,
    cal,
    member_with_access,
    org_base,
    private_root,
    zug,
    zug2,
)
from .test_m13_rsvp import respond, rsvp_event

pytestmark = pytest.mark.django_db
TZ = timezone.get_current_timezone()


def local(y, m, d, h=0, mi=0):
    return timezone.make_aware(datetime.datetime(y, m, d, h, mi), TZ)


class TestTeilnehmer:
    def test_mem_e01_zugexport_summiert_nur_den_zug(self, cal, board_a, zug, zug2, api):
        from events import participants

        _, u1, a1 = member_with_access(cal["a"], "Erika", depts=[zug])
        _, u2, a2 = member_with_access(cal["a"], "Hans", depts=[zug2])
        data = rsvp_event(api(board_a), org_base(cal["a"]), rsvp_guests=True)
        respond(api(u1), a1, data["id"], guests=1)
        respond(api(u2), a2, data["id"], guests=2)
        event = Event.objects.get(pk=data["id"])
        answered, open_ = participants.rows(event, zug.pk)
        sheet = load_workbook(io.BytesIO(participants.xlsx(event, answered, open_))).active
        texts = [row[0].value for row in sheet.iter_rows() if row[0].value]
        assert "Zusagen: 1, Begleitpersonen: 1, Personen: 2" in texts

    def test_mem_e02_rueckmeldung_ausgetretener_entfernen(self, cal, board_a, api):
        m, user, access = member_with_access(cal["a"], "Erika")
        data = rsvp_event(api(board_a), org_base(cal["a"]))
        respond(api(user), access, data["id"])
        MembershipPeriod.objects.filter(membership=m).update(
            exit_date=timezone.localdate() - datetime.timedelta(days=1), exit_reason="exit"
        )
        base = f"{org_base(cal['a'])}events/{data['id']}/responses/"
        assert api(board_a).get(base).json()["summary"]["yes"] == 1
        assert api(board_a).delete(f"{base}{m.pk}/").status_code == 204
        assert not EventResponse.objects.exists()


class TestSerien:
    def _series(self, org, start, end, **extra):
        event = Event.objects.create(
            organization=org, title="Übung", starts_at=start, ends_at=end, **extra
        )
        return Event.objects.get(pk=event.pk)  # Zeiten wie aus der Datenbank (UTC)

    def test_mem_e03_lange_serie_bleibt_sichtbar(self, cal):
        event = self._series(
            cal["a"], local(2000, 1, 4, 19), local(2000, 1, 4, 21), repeat="weekly"
        )
        found = recurrence.occurrences([event], local(2026, 6, 1), local(2026, 6, 30))
        assert len(found) == 4
        assert recurrence.is_rule_start(event, local(2026, 6, 2, 19))

    def test_mem_e04_zeitumstellung_aendert_das_ende_nicht(self, cal):
        # Beginn vor, Ende nach der Umstellung am 26.10.2025 (03:00 → 02:00)
        event = self._series(
            cal["a"], local(2025, 10, 26, 1), local(2025, 10, 26, 4), repeat="yearly"
        )
        first = recurrence.occurrences([event], local(2025, 10, 25), local(2025, 10, 27))[0]
        assert timezone.localtime(first.ends_at).hour == 4

    def test_mem_e08_absage_wandert_mit_der_uhrzeit(self, cal, board_a, api):
        event = self._series(
            cal["a"], local(2026, 3, 3, 19), local(2026, 3, 3, 21), repeat="weekly"
        )
        EventException.objects.create(
            event=event, original_start=local(2026, 4, 7, 19), is_cancelled=True
        )
        resp = api(board_a).patch(
            f"{org_base(cal['a'])}events/{event.pk}/",
            {"starts_at": "2026-03-03T20:00:00", "ends_at": "2026-03-03T22:00:00"},
            content_type=JSON,
        )
        assert resp.status_code == 200, resp.json()
        event.refresh_from_db()
        found = recurrence.occurrences([event], local(2026, 4, 6), local(2026, 4, 9))
        assert [o.is_cancelled for o in found] == [True]
        assert timezone.localtime(found[0].starts_at).hour == 20


class TestStatistik:
    def test_mem_e05_umgekehrter_zeitraum_ist_fehler(self, api, orgs, board_a):
        resp = api(board_a).get(
            f"/api/orgs/{orgs['a'].pk}/members/statistics/?von=2024-12-01&bis=2024-01-31"
        )
        assert resp.status_code == 400

    def test_mem_e06_stichtag_der_altersstruktur(self, api, orgs, board_a):
        resp = api(board_a).get(
            f"/api/orgs/{orgs['a'].pk}/members/statistics/?von=2024-01-01&bis=2024-12-31"
        )
        assert resp.status_code == 200, resp.json()
        assert resp.json()["ages"]["start_on"] == "2023-12-31"

    def test_mem_e07_extremes_eintrittsdatum(self, api, orgs, board_a):
        from .test_m2_members import admit, make_types

        types = make_types(orgs["a"])
        resp = admit(api(board_a), orgs["a"], types["aktiv"], entry_date="9999-12-30")
        assert resp.status_code == 400
        assert "entry_date" in resp.json()
        assert jubilee_date([(datetime.date(9999, 12, 30), None)], 1) is None


class TestZusatzfelder:
    def test_ds04_langer_text_wird_nicht_still_gekuerzt(self):
        assert clean_scalar("text", "x" * 1000) == "x" * 1000
        with pytest.raises(ValidationError):
            clean_scalar("text", "x" * 1001)

    def test_ds05_zahlen(self):
        assert clean_scalar("number", "43,5") == 43.5
        assert clean_scalar("number", "1e3") == 1000
        assert clean_scalar("number", 0) == 0
        for bad in ("NaN", "Infinity", "-inf", "1234567890.1234567", True):
            with pytest.raises(ValidationError):
                clean_scalar("number", bad)

    def test_ds10_null_und_nein_im_listenexport(self, orgs):
        field = FieldDefinition(
            organization=orgs["a"],
            label="Schießen",
            repeatable=True,
            subfields=[
                {"key": "ringe", "label": "Ringe", "type": "number"},
                {"key": "wertung", "label": "Wertung", "type": "bool"},
            ],
        )
        assert format_field(field, [{"ringe": 0, "wertung": False}]) == "0, nein"


class TestImport:
    def test_ds07_teilkontakt_frueherer_mitglieder(self, api, orgs, types, board_a, awards):
        c = api(board_a)
        org = orgs["a"]
        rows = base_rows()
        rows[4][3] = "anton@example.org"  # früheres Mitglied nur mit E-Mail
        job = upload(c, org, rows).json()
        configure(c, org, job, types["a"], awards)
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/")
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/")
        assert ContactDetails.objects.get(person__last_name="Alt").email == "anton@example.org"

    def test_ds09_doppelte_zustellung_aendert_das_journal_nicht(
        self, api, orgs, types, board_a, awards
    ):
        c = api(board_a)
        org = orgs["a"]
        job = upload(c, org, base_rows()).json()
        configure(c, org, job, types["a"], awards)
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/validate/")
        c.post(f"/api/orgs/{org.pk}/imports/{job['id']}/run/")
        before = ImportJob.objects.get(pk=job["id"])
        assert before.status == "imported" and before.created["memberships"]
        tasks.run_import(job["id"])
        after = ImportJob.objects.get(pk=job["id"])
        assert after.status == "imported"
        assert after.created == before.created
        assert after.counts == before.counts


class TestNachverarbeitung:
    def test_k03_fehler_im_aenderungspaket_meldet_keinen_500(self, api, orgs, board_a):
        member = make_member(orgs["a"], "Erika", "Muster")
        url = f"/api/orgs/{orgs['a'].pk}/members/{member.pk}/"
        with mock.patch("changes.middleware.finalize", side_effect=RuntimeError("kaputt")):
            resp = api(board_a).patch(url, {"city": "Neustadt"}, content_type=JSON)
        assert resp.status_code == 200
        assert "X-Change-Set" not in resp
