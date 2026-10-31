# ruff: noqa: F811  (Fixtures aus anderen Testmodulen)
"""Regressionstests zum Audit vom 29.10.2026: Finanzen, Mitglieder, Termine und Kern
(F29-01 bis F29-03, R29-01, MEM-E04/N01/N02/N04, N-K01/N-K02, S29-01, U-K03, R05)."""

import datetime
import io
import uuid
from decimal import Decimal
from types import SimpleNamespace
from unittest import mock

import pytest
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook

from billing.models import Invoice, InvoiceLine, Payment
from billing.views import InvoiceViewSet
from core.tenancy import tenant_context
from customfields.models import FieldDefinition, FieldValue, clean_scalar
from events import participants, recurrence
from events.models import Event
from members.gender import parse_gender, take_over_field
from members.models import Person
from organizations.models import Department, Organization
from organizations.serializers import OrganizationSerializer
from organizations.views import OrganizationDetailView

from .test_audit_finanzen import invoice, pay  # noqa: F401
from .test_m2_members import admit, make_types
from .test_m9_billing import _private, base, c, club, members, setup, start_run, zug  # noqa: F401
from .test_m12_events import cal, member_with_access, org_base, private_root  # noqa: F401
from .test_m13_rsvp import respond, rsvp_event

pytestmark = pytest.mark.django_db
JSON = "application/json"
TZ = timezone.get_current_timezone()


def local(y, m, d, h=0, mi=0, fold=0):
    return datetime.datetime(y, m, d, h, mi, tzinfo=TZ, fold=fold)


class TestSammelrechnung:
    """F29-01: Eine wiederholte Änderung erhält die Beitragsherkunft."""

    def _zug(self, members):
        return Invoice.objects.get(recipient_kind="department")

    def _payload(self, invoice):
        lines = list(invoice.lines.order_by("position"))
        return {
            "lines": [
                {"id": line.pk, "text": line.text, "amount": str(line.amount)} for line in lines
            ]
            + [{"text": "Fahnengeld", "amount": "10.00"}]
        }

    def test_f29_01_wiederholte_aenderung_rechnet_nicht_doppelt(self, c, base, members):
        run = start_run(c, base).json()
        zug = self._zug(members)
        before = {(x.membership_id, x.is_fee) for x in zug.lines.all()}
        data = self._payload(zug)
        key = str(uuid.uuid4())
        url = f"{base}/invoices/{zug.pk}/"
        first = c.patch(url, data, content_type=JSON, HTTP_IDEMPOTENCY_KEY=key)
        again = c.patch(url, data, content_type=JSON, HTTP_IDEMPOTENCY_KEY=key)
        assert first.status_code == again.status_code == 200, again.json()
        lines = list(InvoiceLine.objects.filter(invoice=zug).order_by("position"))
        assert len(lines) == 3  # Fahnengeld nur einmal
        assert {(x.membership_id, x.is_fee) for x in lines[:2]} == before
        assert zug.lines.filter(is_fee=True).count() == 2
        # Ohne Schlüssel (alter Client): die IDs bleiben gleich, die Herkunft auch
        data2 = {"lines": [{**line, "amount": line["amount"]} for line in data["lines"][:2]]}
        assert c.patch(url, data2, content_type=JSON).status_code == 200
        assert c.patch(url, data2, content_type=JSON).status_code == 200
        assert zug.lines.filter(is_fee=True).count() == 2
        c.post(f"{base}/invoices/{zug.pk}/finalize/")
        c.post(f"{base}/runs/{run['id']}/rebuild/")
        assert Invoice.objects.filter(recipient_kind="department").count() == 1
        assert Invoice.objects.filter(recipient_kind="department").get().total == Decimal("90.00")

    def test_f29_01_veraltete_posten_id_wird_abgewiesen(self, c, base, members):
        start_run(c, base)
        zug = self._zug(members)
        data = self._payload(zug)
        data["lines"][0]["id"] = 999999
        resp = c.patch(f"{base}/invoices/{zug.pk}/", data, content_type=JSON)
        assert resp.status_code == 400 and "neu laden" in str(resp.json())
        assert zug.lines.filter(is_fee=True).count() == 2


class TestRechnungsregeln:
    def test_f29_03_faelligkeit_vor_rechnungsdatum_auch_festgeschrieben(self, c, base, invoice):
        resp = c.patch(
            f"{base}/invoices/{invoice.pk}/", {"due_date": "2025-01-01"}, content_type=JSON
        )
        assert resp.status_code == 400
        invoice.refresh_from_db()
        assert invoice.due_date == datetime.date(2025, 2, 15)

    def test_r29_01_hinweis_mit_veraltetem_stand_setzt_zahlung_nicht_zurueck(
        self, c, base, setup, invoice
    ):
        stale = Invoice.objects.get(pk=invoice.pk)  # vor der Zahlung geladen
        assert pay(c, base, invoice, setup).status_code == 200
        with mock.patch.object(InvoiceViewSet, "get_object", return_value=stale):
            resp = c.patch(f"{base}/invoices/{invoice.pk}/", {"note": "Danke"}, content_type=JSON)
        assert resp.status_code == 200, resp.json()
        invoice.refresh_from_db()
        assert invoice.status == Invoice.Status.PAID
        assert invoice.paid == Decimal("60.00") and invoice.note == "Danke"

    def test_u_k03_gleicher_schluessel_anderer_betrag_ist_konflikt(self, c, base, setup, invoice):
        key = str(uuid.uuid4())
        assert pay(c, base, invoice, setup, key=key, amount="20.00").status_code == 200
        resp = pay(c, base, invoice, setup, key=key, amount="40.00")
        assert resp.status_code == 409
        assert Payment.objects.count() == 1


class TestTeilnehmerUntereinheiten:
    def test_mem_n01_gruppe_unter_dem_zug_erscheint_in_der_zugliste(self, cal, board_a, api):
        zug = Department.objects.create(organization=cal["a"], name="Jägerzug")
        gruppe = Department.objects.create(organization=cal["a"], parent=zug, name="Gruppe 1")
        _, user, access = member_with_access(cal["a"], "Erika", depts=[gruppe])
        data = rsvp_event(api(board_a), org_base(cal["a"]), rsvp_guests=True)
        respond(api(user), access, data["id"], guests=2)
        event = Event.objects.get(pk=data["id"])
        answered, open_ = participants.rows(event, zug.pk)
        assert len(answered) == 1
        sheet = load_workbook(io.BytesIO(participants.xlsx(event, answered, open_))).active
        texts = [row[0].value for row in sheet.iter_rows() if row[0].value]
        assert "Zusagen: 1, Begleitpersonen: 2, Personen: 3" in texts


class TestSerienZeitumstellung:
    def test_mem_e04_zweite_herbststunde_behaelt_das_ende(self, cal):
        start = local(2026, 10, 25, 1, 30)  # 01:30 Sommerzeit
        end = local(2026, 10, 25, 2, 30, fold=1)  # zweites 02:30 (Winterzeit)
        event = Event.objects.create(
            organization=cal["a"],
            title="Nachtübung",
            starts_at=start,
            ends_at=end,
            repeat="yearly",
        )
        event = Event.objects.get(pk=event.pk)
        first = recurrence.occurrences([event], local(2026, 10, 24), local(2026, 10, 26))[0]
        assert first.ends_at == event.ends_at
        assert first.ends_at - first.starts_at == datetime.timedelta(hours=2)
        # Spätere Termine enden nie vor ihrem Beginn
        later = recurrence.occurrences([event], local(2027, 10, 24), local(2027, 10, 26))
        assert later and later[0].ends_at > later[0].starts_at


class TestGeschlecht:
    @pytest.mark.parametrize("value", ["k.a", "k. a", "K. A.", "keine Angabe"])
    def test_mem_n04_dokumentierte_leere_angaben(self, value):
        assert parse_gender(value) == ""

    def test_mem_n02_widerspruch_bleibt_leer_und_zaehlt_einmal(self, orgs):
        defs = [
            FieldDefinition.objects.create(
                organization=orgs[key], label="Geschlecht", key="geschlecht", field_type="text"
            )
            for key in ("a", "b")
        ]
        with tenant_context(orgs["a"].pk):
            both = Person.objects.create(
                scope=orgs["federation"],
                last_name="Zwei",
                first_name="Werte",
                birth_date="1990-01-01",
            )
            same = Person.objects.create(
                scope=orgs["federation"],
                last_name="Gleich",
                first_name="Werte",
                birth_date="1990-01-01",
            )
        for definition, a, b in zip(defs, ("Mann", "Frau"), ("Frau", "weiblich"), strict=True):
            org = definition.organization
            FieldValue.objects.create(organization=org, definition=definition, person=both, value=a)
            FieldValue.objects.create(organization=org, definition=definition, person=same, value=b)
        assert take_over_field(Person, FieldDefinition, FieldValue) == 1
        assert Person.objects.get(pk=both.pk).gender == ""  # Widerspruch: von Hand klären
        assert Person.objects.get(pk=same.pk).gender == "weiblich"


class TestKern:
    def test_n_k01_zwei_gueltige_aenderungen_ergeben_keinen_30_februar(self, orgs, board_a):
        org = orgs["a"]
        stale = Organization.objects.get(pk=org.pk)
        ser = OrganizationSerializer(
            stale, data={"fiscal_start_day": 30}, partial=True, context={"org": org}
        )
        assert ser.is_valid(), ser.errors  # gegen Januar gültig
        # In der Zwischenzeit stellt jemand den Monat auf Februar (Tag 1: gültig)
        Organization.objects.filter(pk=org.pk).update(fiscal_start_month=2)
        view = OrganizationDetailView()
        view.org = org
        view.request = SimpleNamespace(user=board_a)
        with pytest.raises(Exception) as exc:
            view.perform_update(ser)
        assert "fiscal_start_day" in str(exc.value.detail)
        org.refresh_from_db()
        assert (org.fiscal_start_month, org.fiscal_start_day) == (2, 1)

    def test_n_k02_fehler_in_der_beschriftung_meldet_keinen_500(self, api, orgs, board_a):
        from .factories import make_member

        member = make_member(orgs["a"], "Erika", "Muster")
        url = f"/api/orgs/{orgs['a'].pk}/members/{member.pk}/"
        with mock.patch("changes.middleware.label", side_effect=RuntimeError("kaputt")):
            resp = api(board_a).patch(url, {"city": "Neustadt"}, content_type=JSON)
        assert resp.status_code == 200
        assert resp["X-Change-Label"] == "%C3%84nderung"

    @pytest.mark.parametrize("value", ["1e-400", "1e-16", "1e15", "1e400"])
    def test_s29_01_zu_grosse_oder_kleine_zahlen_werden_abgelehnt(self, value):
        with pytest.raises(ValidationError):
            clean_scalar("number", value)

    @pytest.mark.parametrize(
        "value,expected", [("0.000000000000001", 1e-15), ("999999999999999", 999999999999999)]
    )
    def test_s29_01_grenzwerte_bleiben_erlaubt(self, value, expected):
        assert clean_scalar("number", value) == expected

    def test_r05_geloeschte_abteilung_bei_aufnahme(self, api, orgs, board_a):
        types = make_types(orgs["a"])
        gone = Department.objects.create(organization=orgs["a"], name="Aufgelöst")
        gone_id = gone.pk
        gone.delete()
        resp = admit(api(board_a), orgs["a"], types["aktiv"], department_ids=[gone_id])
        assert resp.status_code == 400 and "department_ids" in resp.json()
        assert not Person.objects.exists()
