"""Wettbewerbe Schritt 4: Auszeichnungen, Titel, Würdenvorschlag, Wanderpokale, Reihen und
Mitgliedsakte (alles erfunden)."""

import pytest

from competitions.models import AwardGrant, Competition
from honors.models import Dignity, Reign

from .test_wettbewerbe import _punkte, _starter, club, schuetzen, url, zuege  # noqa: F401


def _schiessen(c, club, title="Vereinsmeisterschaft 2027", date="2027-06-14", series=None):  # noqa: F811
    data = {"title": title, "date": date, "template": "schiessen"}
    if series:
        data["series"] = series
    resp = c.post(url(club), data, "application/json")
    assert resp.status_code == 201, resp.content
    comp_id = resp.json()["id"]
    einzel, zug = c.get(url(club, f"{comp_id}/")).json()["scorings"]
    c.patch(url(club, f"{comp_id}/scorings/{zug['id']}/"), {"team_size": 2}, "application/json")
    return comp_id, einzel["id"], zug["id"]


def _final(c, club, comp_id):  # noqa: F811
    resp = c.post(url(club, f"{comp_id}/status/"), {"status": "endgueltig"}, "application/json")
    assert resp.status_code == 200, resp.content


@pytest.mark.django_db(transaction=True)
class TestAuszeichnungen:
    def test_titel_wuerde_pokal_und_akte(self, club, schuetzen, zuege, board_a, api):  # noqa: F811
        c = api(board_a)
        koenig = Dignity.objects.create(organization=club, name="Schützenkönig")
        reihe = c.post(url(club, "series/"), {"name": "Vereinsmeisterschaft"}, "application/json")
        assert reihe.status_code == 201, reihe.content
        reihe_id = reihe.json()["id"]
        pokal = c.post(
            url(club, "trophies/"), {"name": "Zugpokal", "keep_after": 2}, "application/json"
        ).json()
        comp_id, einzel, zug = _schiessen(c, club, series=reihe_id)
        r = c.post(
            url(club, f"{comp_id}/awards/"),
            {
                "scoring": einzel,
                "place_from": 1,
                "place_to": 1,
                "title": "Vereinsmeister",
                "dignity": koenig.pk,
            },
            "application/json",
        )
        assert r.status_code == 201, r.content
        r = c.post(
            url(club, f"{comp_id}/awards/"),
            {"scoring": einzel, "place_from": 1, "place_to": 3, "item": "Medaille"},
            "application/json",
        )
        assert r.status_code == 201
        r = c.post(
            url(club, f"{comp_id}/awards/"),
            {"scoring": zug, "place_from": 1, "place_to": 1, "trophy": pokal["id"]},
            "application/json",
        )
        assert r.status_code == 201
        # ohne Inhalt geht nicht
        bad = c.post(url(club, f"{comp_id}/awards/"), {"scoring": einzel}, "application/json")
        assert bad.status_code == 400
        st = {f: _starter(c, club, comp_id, membership=m.pk) for f, m in schuetzen.items()}
        punkte = {
            "Max": 30,
            "Erika": 27,
            "Otto": 20,
            "Jürgen": 29,
            "Anna": 26,
            "Paul": 25,
            "Lena": 22,
        }
        for f, p in punkte.items():
            assert _punkte(c, club, comp_id, einzel, st[f]["id"], p).status_code == 200
        _final(c, club, comp_id)
        grants = c.get(url(club, f"{comp_id}/")).json()["grants"]
        titel = [g for g in grants if g["title"]]
        assert [(g["full_title"], g["name"]) for g in titel] == [
            ("Vereinsmeister 2027", "Max Muster")
        ]
        assert titel[0]["dignity_state"] == "vorschlag"
        assert sorted(g["name"] for g in grants if g["item"] == "Medaille") == [
            "Erika Beispiel",
            "Jürgen Müller",
            "Max Muster",
        ]
        # Zug 1: 30 + 27 = 57, Zug 2: 29 + 26 = 55
        pokal_g = [g for g in grants if g["trophy"]]
        assert [g["name"] for g in pokal_g] == ["Zug 1"]
        # Würde bestätigen: legt die Regentschaft mit Herkunft Zug an
        vorschlag = c.get(url(club, "grants/?vorschlaege=1")).json()
        assert len(vorschlag) == 1
        r = c.post(
            url(club, f"grants/{vorschlag[0]['id']}/dignity/"),
            {"decision": "bestaetigen"},
            "application/json",
        )
        assert r.status_code == 200, r.content
        reign = Reign.objects.get()
        assert reign.year == 2027 and reign.dignity == koenig
        assert reign.person == schuetzen["Max"].person and reign.origin_id == zuege[1].pk
        again = c.post(
            url(club, f"grants/{vorschlag[0]['id']}/dignity/"),
            {"decision": "bestaetigen"},
            "application/json",
        )
        assert again.status_code == 400
        # Korrektur: offene Vergaben fallen weg, die bestätigte Würde bleibt
        c.post(url(club, f"{comp_id}/status/"), {"status": "vorlaeufig"}, "application/json")
        left = AwardGrant.objects.filter(competition_id=comp_id)
        assert [g.dignity_state for g in left] == ["bestaetigt"]
        _final(c, club, comp_id)
        assert AwardGrant.objects.filter(competition_id=comp_id).count() == 5
        assert Reign.objects.count() == 1
        # Mitgliedsakte
        akte = c.get(url(club, f"member/{schuetzen['Max'].pk}/")).json()
        assert akte["starts"][0]["results"][0]["place"] == 1
        assert {g["full_title"] for g in akte["grants"]} == {"Vereinsmeister 2027", "Medaille"}

        # Folgejahr aus der Kopie: Reihe und Auszeichnungen kommen mit; Zug 1 gewinnt erneut
        copy = c.post(url(club, f"{comp_id}/copy/"), {}, "application/json").json()
        detail = c.get(url(club, f"{copy['id']}/")).json()
        assert detail["series"] == reihe_id and len(detail["awards"]) == 3
        e2, z2 = (s["id"] for s in detail["scorings"])
        st2 = {f: _starter(c, club, copy["id"], membership=m.pk) for f, m in schuetzen.items()}
        for f, p in punkte.items():
            _punkte(c, club, copy["id"], e2, st2[f]["id"], p)
        _final(c, club, copy["id"])
        holders = c.get(url(club, f"trophies/{pokal['id']}/")).json()["holders"]
        assert [(h["year"], h["name"], h["streak"], h["kept"]) for h in holders] == [
            (2028, "Zug 1", 2, True),
            (2027, "Zug 1", 1, False),
        ]
        # Altdaten von Hand
        r = c.post(
            url(club, "grants/"),
            {"year": 1999, "title": "Vereinsmeister", "name": "Karl Früher", "series": reihe_id},
            "application/json",
        )
        assert r.status_code == 201, r.content
        archiv = c.get(url(club, f"series/{reihe_id}/")).json()
        assert [y["year"] for y in archiv["years"]] == [2028, 2027]
        assert archiv["years"][0]["winners"][0]["podium"][0]["name"] == "Max Muster"
        assert {(t["year"], t["name"]) for t in archiv["titles"] if t["title"]} >= {
            (1999, "Karl Früher"),
            (2027, "Max Muster"),
            (2028, "Max Muster"),
        }
        rekorde = {r["scoring"]: r for r in archiv["records"]}
        assert rekorde["Einzelwertung"]["value_text"] == "30"
        assert rekorde["Einzelwertung"]["name"] == "Max Muster"
        assert archiv["units"][0]["name"] == "Zug 1"
        assert [y["place"] for y in archiv["units"][0]["years"]] == [1, 1]

    def test_wuerde_ablehnen_und_rechte(self, club, schuetzen, board_a, api):  # noqa: F811
        from tests.factories import give_role, make_user

        c = api(board_a)
        koenig = Dignity.objects.create(organization=club, name="König")
        comp_id, einzel, _ = _schiessen(c, club)
        c.post(
            url(club, f"{comp_id}/awards/"),
            {"scoring": einzel, "dignity": koenig.pk},
            "application/json",
        )
        s = _starter(c, club, comp_id, membership=schuetzen["Otto"].pk)
        _punkte(c, club, comp_id, einzel, s["id"], 25)
        _final(c, club, comp_id)
        g = AwardGrant.objects.get()
        # Schießmeister ohne Mitgliederrecht darf nicht bestätigen
        meister = make_user("meister@example.org")
        give_role(meister, club, permissions=["competitions.manage"], name="Schießmeister")
        r = api(meister).post(
            url(club, f"grants/{g.pk}/dignity/"), {"decision": "bestaetigen"}, "application/json"
        )
        assert r.status_code == 403
        r = c.post(
            url(club, f"grants/{g.pk}/dignity/"), {"decision": "ablehnen"}, "application/json"
        )
        assert r.status_code == 200 and r.json()["dignity_state"] == "abgelehnt"
        assert not Reign.objects.exists()
        # Vergabe aus dem Wettbewerb lässt sich nicht direkt löschen
        assert c.delete(url(club, f"grants/{g.pk}/")).status_code == 400
        assert Competition.objects.get(pk=comp_id).status == "endgueltig"

    def test_pfand_mit_wuerde(self, club, schuetzen, board_a, api):  # noqa: F811
        c = api(board_a)
        koenig = Dignity.objects.create(organization=club, name="König")
        resp = c.post(
            url(club),
            {"title": "Vogelschießen 2027", "date": "2027-05-01", "template": "vogel"},
            "application/json",
        )
        comp_id = resp.json()["id"]
        detail = c.get(url(club, f"{comp_id}/")).json()
        rumpf = detail["prizes"][-1]
        r = c.post(
            url(club, f"{comp_id}/awards/"),
            {"prize": rumpf["id"], "title": "Schützenkönig", "dignity": koenig.pk},
            "application/json",
        )
        assert r.status_code == 201, r.content
        s = _starter(c, club, comp_id, membership=schuetzen["Anna"].pk)
        r = c.patch(
            url(club, f"{comp_id}/prizes/{rumpf['id']}/"),
            {"winner": s["id"], "shot_number": 143},
            "application/json",
        )
        assert r.status_code == 200, r.content
        _final(c, club, comp_id)
        g = AwardGrant.objects.get()
        assert g.name == "Anna Schmidt" and g.dignity_state == "vorschlag"
        assert g.full_title == "Schützenkönig 2027"

    def test_export_spalte(self, settings, club, schuetzen, board_a, api):  # noqa: F811
        import io

        from openpyxl import load_workbook

        settings.CELERY_TASK_ALWAYS_EAGER = True
        AwardGrant.objects.create(
            organization=club,
            year=2019,
            title="Vereinsmeister",
            name="Max Muster",
            membership=schuetzen["Max"],
            manual=True,
        )
        AwardGrant.objects.create(
            organization=club,
            year=2023,
            title="Vereinsmeister",
            name="Max Muster",
            membership=schuetzen["Max"],
            manual=True,
        )
        c = api(board_a)
        resp = c.post(
            f"/api/orgs/{club.pk}/exports/",
            {
                "kind": "members",
                "file_format": "xlsx",
                "columns": ["last_name", "competition_titles"],
            },
            content_type="application/json",
        )
        assert resp.status_code == 201, resp.json()
        file = c.get(f"/api/orgs/{club.pk}/exports/{resp.json()['id']}/download/")
        ws = load_workbook(io.BytesIO(b"".join(file.streaming_content))).active
        values = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
        assert values["Muster"] == "Vereinsmeister: 2019, 2023"
