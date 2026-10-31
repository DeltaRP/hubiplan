"""Würden und Chronik: Verwaltung, Excel-Import, öffentliche Sicht (alles erfunden)."""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import Workbook

from honors.models import Companion, Dignity, Reign, ReignPhoto
from organizations.models import Department
from website import services
from website.models import Website

from .factories import give_role, make_member, make_user
from .test_webseite_inhalte import png

pytestmark = pytest.mark.django_db(transaction=True)


@pytest.fixture(autouse=True)
def private_root(settings, tmp_path):
    settings.PRIVATE_ROOT = tmp_path / "private"


def url(org, rest=""):
    return f"/api/orgs/{org.pk}/honors/{rest}"


@pytest.fixture
def people(orgs):
    org = orgs["a"]
    return {
        "erika": make_member(org, "Erika", "Beispiel"),
        "otto": make_member(org, "Otto", "Muster"),
        "jürgen": make_member(org, "Jürgen", "Schütz"),
    }


@pytest.fixture
def koenig(orgs):
    return Dignity.objects.create(organization=orgs["a"], name="König", sort_order=1)


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("koenige.xlsx", buf.getvalue())


class TestVerwaltung:
    def test_wuerde_anlegen_und_doppelt(self, orgs, board_a, api):
        c = api(board_a)
        resp = c.post(url(orgs["a"], "dignities/"), {"name": "Schülerprinz"}, "application/json")
        assert resp.status_code == 201, resp.content
        resp = c.post(url(orgs["a"], "dignities/"), {"name": "schülerprinz"}, "application/json")
        assert resp.status_code == 400

    def test_wuerde_einer_untereinheit(self, orgs, board_a, board_b, api):
        unit = Department.objects.create(organization=orgs["a"], parent=orgs["a"], name="Zug 1")
        other = Department.objects.create(organization=orgs["b"], parent=orgs["b"], name="Zug B")
        c = api(board_a)
        ok = c.post(
            url(orgs["a"], "dignities/"), {"name": "Zugkönig", "unit": unit.pk}, "application/json"
        )
        assert ok.status_code == 201 and ok.json()["unit_name"] == "Zug 1"
        bad = c.post(
            url(orgs["a"], "dignities/"), {"name": "Zugkönig", "unit": other.pk}, "application/json"
        )
        assert bad.status_code == 400
        assert api(board_b).get(url(orgs["a"], "dignities/")).status_code == 404

    def test_regentschaft_mit_mitglied_und_begleitung(self, orgs, board_a, api, people, koenig):
        c = api(board_a)
        resp = c.post(
            url(orgs["a"], "reigns/"),
            {
                "dignity": koenig.pk,
                "year": 2024,
                "person": people["otto"].person_id,
                "companions": [
                    {"label": "Königin", "person": people["erika"].person_id},
                    {"label": "Minister", "name": "Karl Früher"},
                ],
            },
            "application/json",
        )
        assert resp.status_code == 201, resp.content
        data = resp.json()
        assert data["display_name"] == "Otto Muster" and data["name"] == "Otto Muster"
        assert [x["display_name"] for x in data["companions"]] == ["Erika Beispiel", "Karl Früher"]
        # Filter nach Person findet auch Begleitpersonen
        found = c.get(url(orgs["a"], f"reigns/?person={people['erika'].person_id}")).json()
        assert len(found) == 1

    def test_regentschaft_braucht_jahr_und_name(self, orgs, board_a, api, koenig):
        c = api(board_a)
        assert (
            c.post(
                url(orgs["a"], "reigns/"), {"dignity": koenig.pk, "name": "X"}, "application/json"
            )
        ).status_code == 400
        assert (
            c.post(
                url(orgs["a"], "reigns/"), {"dignity": koenig.pk, "year": 1950}, "application/json"
            )
        ).status_code == 400
        ok = c.post(
            url(orgs["a"], "reigns/"),
            {"dignity": koenig.pk, "year": 1950, "name": "Heinrich Vorzeit"},
            "application/json",
        )
        assert ok.status_code == 201

    def test_regentschaft_mit_startjahr_und_endjahr(self, orgs, board_a, api, koenig):
        """Ralf 2026-10-30: Startjahr/Endjahr ohne genaues Datum."""
        c = api(board_a)
        reigns = url(orgs["a"], "reigns/")
        body = {"dignity": koenig.pk, "year": 1950, "end_year": 1952, "name": "Heinrich Vorzeit"}
        ok = c.post(reigns, body, "application/json")
        assert ok.status_code == 201
        assert ok.json()["end_year"] == 1952
        assert ok.json()["period_text"] == "1950–1952"
        # Endjahr vor dem Startjahr oder ohne Startjahr: abgelehnt
        bad = c.post(reigns, {**body, "end_year": 1949}, "application/json")
        assert bad.status_code == 400
        assert "end_year" in bad.json()
        no_start = {"dignity": koenig.pk, "end_year": 1952, "name": "X"}
        assert c.post(reigns, no_start, "application/json").status_code == 400

    def test_fremdes_mitglied_abgelehnt(self, orgs, board_a, api, koenig):
        stranger = make_member(orgs["b"], "Fremd", "Person")
        resp = api(board_a).post(
            url(orgs["a"], "reigns/"),
            {"dignity": koenig.pk, "year": 2020, "person": stranger.person_id},
            "application/json",
        )
        assert resp.status_code == 400

    def test_nur_lesen_ohne_bearbeitungsrecht(self, orgs, api, koenig):
        user = make_user("leser@example.org")
        give_role(user, orgs["a"], permissions=["members.view"], name="Leser")
        c = api(user)
        assert c.get(url(orgs["a"], "reigns/")).status_code == 200
        resp = c.post(
            url(orgs["a"], "reigns/"),
            {"dignity": koenig.pk, "year": 2000, "name": "A B"},
            "application/json",
        )
        assert resp.status_code == 403

    def test_wuerde_mit_regentschaft_nur_mit_bestaetigung(self, orgs, board_a, api, koenig):
        reign = Reign.objects.create(organization=orgs["a"], dignity=koenig, year=1990, name="A B")
        Companion.objects.create(organization=orgs["a"], reign=reign, label="Königin", name="C D")
        other = Dignity.objects.create(organization=orgs["a"], name="Kaiser")
        Reign.objects.create(organization=orgs["a"], dignity=other, year=1991, name="E F")
        c = api(board_a)
        assert c.delete(url(orgs["a"], f"dignities/{koenig.pk}/")).status_code == 400
        assert Dignity.objects.filter(pk=koenig.pk).exists()
        resp = c.delete(url(orgs["a"], f"dignities/{koenig.pk}/") + "?mit_regentschaften=1")
        assert resp.status_code == 204, resp.content
        assert not Dignity.objects.filter(pk=koenig.pk).exists()
        assert not Companion.objects.exists()
        # andere Würden bleiben unberührt
        assert list(Reign.objects.values_list("name", flat=True)) == ["E F"]

    def test_foto_hochladen(self, orgs, board_a, api, koenig):
        reign = Reign.objects.create(organization=orgs["a"], dignity=koenig, year=1990, name="A B")
        c = api(board_a)
        resp = c.post(url(orgs["a"], f"reigns/{reign.pk}/photo/"), {"file": png()})
        assert resp.status_code == 201, resp.content
        photo = ReignPhoto.objects.get(reign=reign)
        assert photo.width == 2000
        assert c.get(url(orgs["a"], f"reigns/{reign.pk}/photo/")).status_code == 200
        assert c.delete(url(orgs["a"], f"reigns/{reign.pk}/photo/")).status_code == 204
        assert not ReignPhoto.objects.exists()

    def test_personensuche(self, orgs, board_a, api, people):
        rows = api(board_a).get(url(orgs["a"], "people/?suche=mus ot")).json()
        assert rows == [
            {
                "id": people["otto"].person_id,
                "first_name": "Otto",
                "last_name": "Muster",
                "birth_year": rows[0]["birth_year"],
            }
        ]


class TestImport:
    def test_vorschau_schlaegt_mitglieder_vor(self, orgs, board_a, api, people, koenig):
        file = xlsx(
            [
                ["Jahr", "Würde", "Name", "Begleitung", "Bemerkung"],
                [2023, "König", "Muster, Otto", "Königin: Erika Beispiel", ""],
                [1961, "König", "Wilhelm Vorzeit", "", "Jubiläumskönig"],
                ["1999/2000", "Prinz", "Juergen Schuetz", "", ""],
                ["", "König", "", "", ""],
            ]
        )
        resp = api(board_a).post(url(orgs["a"], "import/preview/"), {"file": file})
        assert resp.status_code == 200, resp.content
        data = resp.json()
        rows = data["rows"]
        assert rows[0]["person_id"] == people["otto"].person_id
        assert rows[0]["companions"][0]["person_id"] == people["erika"].person_id
        assert rows[0]["dignity_id"] == koenig.pk
        assert rows[1]["person_id"] is None and rows[1]["note"] == "Jubiläumskönig"
        # Umlaute anders geschrieben: trotzdem erkannt
        assert rows[2]["year"] == 1999 and rows[2]["person_id"] == people["jürgen"].person_id
        assert rows[3]["errors"]
        assert data["new_dignities"] == ["Prinz"]
        assert not Reign.objects.exists()  # Vorschau speichert nichts

    def test_fehlende_spalten(self, orgs, board_a, api):
        file = xlsx([["Jahr", "Name"], [2000, "A B"]])
        resp = api(board_a).post(url(orgs["a"], "import/preview/"), {"file": file})
        assert resp.status_code == 400 and "Würde" in str(resp.json())

    def test_uebernehmen(self, orgs, board_a, api, people, koenig):
        resp = api(board_a).post(
            url(orgs["a"], "import/"),
            {
                "rows": [
                    {
                        "year": 2023,
                        "dignity": "König",
                        "name": "Muster, Otto",
                        "person_id": people["otto"].person_id,
                        "companions": [{"label": "Königin", "name": "Erika Beispiel"}],
                    },
                    {"year": 1961, "dignity": "Kaiser", "name": "Wilhelm Vorzeit"},
                ]
            },
            "application/json",
        )
        assert resp.status_code == 201, resp.content
        assert resp.json() == {"created": 2, "dignities_created": ["Kaiser"]}
        reign = Reign.objects.get(year=2023)
        assert reign.dignity == koenig and reign.name == "Otto Muster"
        assert Companion.objects.get().name == "Erika Beispiel"

    def test_ordnungszahl_aus_alter_liste(self, orgs, board_a, api, koenig):
        """Ralf 2026-10-30: Zahlen einer alten Königsliste unverändert übernehmen."""
        file = xlsx(
            [
                ["Jahr", "Würde", "Name", "Ordnungszahl"],
                [1936, "König", "Hubert Beispiel", "IV"],
                [1949, "König", "Hubert Beispiel", "V."],
                [1950, "König", "Otto Vorzeit", 3],
                [1951, "König", "Erika Früh", ""],
                [1952, "König", "Max Muster", "IIII"],
                [1953, "König", "Max Muster", "vier"],
            ]
        )
        rows = api(board_a).post(url(orgs["a"], "import/preview/"), {"file": file}).json()["rows"]
        assert [r["ordinal"] for r in rows] == [4, 5, 3, None, None, None]
        assert rows[1]["ordinal_text"] == "V"
        assert [bool(r["errors"]) for r in rows] == [False] * 4 + [True, True]
        resp = api(board_a).post(
            url(orgs["a"], "import/"),
            {
                "rows": [
                    {
                        "year": r["year"],
                        "dignity": "König",
                        "name": r["name"],
                        "ordinal": r["ordinal"],
                    }
                    for r in rows[:4]
                ]
            },
            "application/json",
        )
        assert resp.status_code == 201, resp.content
        koenig.numbered = True
        koenig.save()
        data = api(board_a).get(url(orgs["a"], "reigns/")).json()
        names = {r["year"]: r["numbered_name"] for r in data}
        # gleicher Name, verschiedene Zahlen: die Zahl aus der Liste gilt je Regentschaft
        assert names == {
            1936: "Hubert IV. Beispiel",
            1949: "Hubert V. Beispiel",
            1950: "Otto III. Vorzeit",
            1951: "Erika I. Früh",
        }

    def test_endjahr(self, orgs, board_a, api, koenig):
        """Ralf 2026-10-30: Würden über mehrere Jahre mit Spalte Endjahr importieren."""
        file = xlsx(
            [
                ["Jahr", "Endjahr", "Würde", "Name"],
                [2019, 2022, "König", "Otto Lang"],
                [2023, "", "König", "Erika Kurz"],
                [2024, 2024, "König", "Max Gleich"],
                [2025, 2020, "König", "Max Falsch"],
                [2026, "bald", "König", "Max Unklar"],
            ]
        )
        rows = api(board_a).post(url(orgs["a"], "import/preview/"), {"file": file}).json()["rows"]
        assert [r["end_year"] for r in rows[:3]] == [2022, None, None]
        assert [bool(r["errors"]) for r in rows] == [False, False, False, True, True]
        resp = api(board_a).post(
            url(orgs["a"], "import/"),
            {
                "rows": [
                    {
                        "year": r["year"],
                        "end_year": r["end_year"],
                        "dignity": "König",
                        "name": r["name"],
                    }
                    for r in rows[:3]
                ]
            },
            "application/json",
        )
        assert resp.status_code == 201, resp.content
        periods = {r.name: r.period_text for r in Reign.objects.all()}
        assert periods == {"Otto Lang": "2019–2022", "Erika Kurz": "2023", "Max Gleich": "2024"}
        bad = api(board_a).post(
            url(orgs["a"], "import/"),
            {"rows": [{"year": 2025, "end_year": 2020, "dignity": "König", "name": "X Y"}]},
            "application/json",
        )
        assert bad.status_code == 400

    def test_uebernehmen_nur_eigene_mitglieder(self, orgs, board_a, api, koenig):
        stranger = make_member(orgs["b"], "Fremd", "Person")
        resp = api(board_a).post(
            url(orgs["a"], "import/"),
            {"rows": [{"year": 2000, "dignity": "König", "person_id": stranger.person_id}]},
            "application/json",
        )
        assert resp.status_code == 400
        assert not Reign.objects.exists()


class TestExport:
    def test_eine_spalte_je_wuerde(self, settings, orgs, board_a, api, people, koenig):
        """Ralf 2026-10-30: im Mitglieder-Export pro definierter Würde eine Spalte."""
        from openpyxl import load_workbook

        settings.CELERY_TASK_ALWAYS_EAGER = True
        org = orgs["a"]
        unit = Department.objects.create(organization=org, parent=org, name="Zug 1")
        zugkoenig = Dignity.objects.create(organization=org, unit=unit, name="Zugkönig")
        Dignity.objects.create(organization=org, name="Alt", is_archived=True)
        otto = people["otto"].person
        Reign.objects.create(organization=org, dignity=koenig, year=2020, person=otto)
        Reign.objects.create(
            organization=org, dignity=koenig, year=2010, end_year=2011, person=otto
        )
        Reign.objects.create(organization=org, dignity=zugkoenig, year=2015, person=otto)
        c = api(board_a)
        cols = c.get(f"/api/orgs/{org.pk}/export-columns/").json()["members"]
        labels = {x["key"]: x["label"] for x in cols}
        assert labels[f"dignity:{koenig.pk}"] == "König"
        assert labels[f"dignity:{zugkoenig.pk}"] == "Zugkönig (Zug 1)"
        assert "Alt" not in labels.values()
        resp = c.post(
            f"/api/orgs/{org.pk}/exports/",
            {
                "kind": "members",
                "file_format": "xlsx",
                "columns": ["last_name", f"dignity:{koenig.pk}", f"dignity:{zugkoenig.pk}"],
            },
            content_type="application/json",
        )
        assert resp.status_code == 201, resp.json()
        file = c.get(f"/api/orgs/{org.pk}/exports/{resp.json()['id']}/download/")
        ws = load_workbook(io.BytesIO(b"".join(file.streaming_content))).active
        values = {row[0]: row[1:] for row in ws.iter_rows(min_row=2, values_only=True)}
        assert values["Muster"] == ("2010–2011, 2020", "2015")
        assert values["Beispiel"] == (None, None)


@pytest.fixture
def site(orgs):
    org = orgs["a"]
    org.enabled_modules = ["members", "website"]
    org.save()
    return Website.objects.create(organization=org, public_key="schluessel-w", state="public")


class TestOeffentlich:
    def test_chronik_mit_widerspruch_und_zaehlung(self, orgs, site, people, koenig):
        org = orgs["a"]
        otto, erika = people["otto"], people["erika"]
        for year in (2019, 2023):
            Reign.objects.create(organization=org, dignity=koenig, year=year, person=otto.person)
        r = Reign.objects.create(organization=org, dignity=koenig, year=2024, person=erika.person)
        Companion.objects.create(organization=org, reign=r, label="Prinzgemahl", name="Max Alt")
        Reign.objects.create(organization=org, dignity=koenig, year=1950, name="Heinrich Vorzeit")
        hidden = Dignity.objects.create(organization=org, name="Intern", is_public=False)
        Reign.objects.create(organization=org, dignity=hidden, year=2024, name="Geheim Name")
        services.set_entry(otto, "honors", True, "paper", None)  # Widerspruch

        resp = Client().get("/api/public/site/v1/schluessel-w/honors/")
        assert resp.status_code == 200, resp.content
        dignities = resp.json()["dignities"]
        assert [d["name"] for d in dignities] == ["König"]
        reigns = dignities[0]["reigns"]
        assert [x["year"] for x in reigns] == [2024, 1950]  # Otto fehlt (Widerspruch)
        assert dignities[0]["current"][0]["holder"]["name"] == "Erika Beispiel"
        assert reigns[0]["companions"] == [
            {"label": "Prinzgemahl", "name": "Max Alt", "photo_url": None}
        ]
        assert "Geheim" not in resp.content.decode()

    def test_traeger_im_laufenden_jahr(self, orgs, site, koenig):
        """Ralf 31.10.: ohne Träger im laufenden Jahr zeigt die Webseite alle."""
        import datetime

        org = orgs["a"]
        Reign.objects.create(organization=org, dignity=koenig, year=2008, end_year=2009, name="A B")

        def flag():
            resp = Client().get("/api/public/site/v1/schluessel-w/honors/")
            return resp.json()["dignities"][0]["current_year"]

        assert flag() is False
        this_year = datetime.date.today().year
        Reign.objects.create(
            organization=org, dignity=koenig, year=this_year - 1, end_year=this_year, name="C D"
        )
        assert flag() is True

    def test_mehrmaliger_koenig_und_nn(self, orgs, site, people, koenig):
        org = orgs["a"]
        for year in (2010, 2015):
            Reign.objects.create(organization=org, dignity=koenig, year=year, name="Paul Doppelt")
        services.set_entry(people["otto"], "honors", True, "paper", None)
        Reign.objects.create(
            organization=org, dignity=koenig, year=2020, person=people["otto"].person
        )
        site.missing_consent = "nn"
        site.save()
        reigns = (
            Client()
            .get("/api/public/site/v1/schluessel-w/honors/")
            .json()["dignities"][0]["reigns"]
        )
        assert reigns[0]["holder"]["name"] == "N. N."
        assert [x["times"] for x in reigns[1:]] == [2, 1]

    def test_foto_nur_mit_einwilligung(self, orgs, site, people, koenig, board_a, api):
        org = orgs["a"]
        old = Reign.objects.create(organization=org, dignity=koenig, year=1960, name="Hans Alt")
        new = Reign.objects.create(
            organization=org, dignity=koenig, year=2024, person=people["otto"].person
        )
        c = api(board_a)
        for reign in (old, new):
            assert c.post(url(org, f"reigns/{reign.pk}/photo/"), {"file": png()}).status_code == 201
        reigns = (
            Client()
            .get("/api/public/site/v1/schluessel-w/honors/")
            .json()["dignities"][0]["reigns"]
        )
        assert reigns[0]["photo_url"] is None  # Otto hat nicht eingewilligt
        assert reigns[1]["photo_url"]
        path = reigns[1]["photo_url"].split("/api/public/", 1)[1]
        assert Client().get(f"/api/public/{path}").status_code == 200
        assert (
            Client().get("/api/public/site/v1/schluessel-w/honors/photos/falsch/").status_code
            == 404
        )

    def test_baustein_aus(self, orgs, site, koenig):
        site.blocks = {**site.blocks, "honors": False}
        site.save()
        assert Client().get("/api/public/site/v1/schluessel-w/honors/").status_code == 404

    def test_wuerden_der_untereinheit(self, orgs, site, koenig):
        org = orgs["a"]
        unit = Department.objects.create(organization=org, parent=org, name="Jägerzug")
        zk = Dignity.objects.create(organization=org, unit=unit, name="Zugkönig")
        Reign.objects.create(organization=org, dignity=zk, year=2022, name="Zug Mann")
        data = Client().get("/api/public/site/v1/schluessel-w/units/jaegerzug/").json()
        assert data["honors"][0]["name"] == "Zugkönig"
        top = Client().get("/api/public/site/v1/schluessel-w/honors/").json()["dignities"]
        assert top == []


class TestOrdnungszahl:
    """Ralf 2026-10-30: römische Zahl je gleichem Vornamen, „Vorname Zahl. Nachname“."""

    def test_roemische_zahlen(self):
        from honors.numbering import roman

        assert [roman(n) for n in (1, 2, 4, 9, 14, 40, 1999)] == [
            "I",
            "II",
            "IV",
            "IX",
            "XIV",
            "XL",
            "MCMXCIX",
        ]

    def test_automatisch_je_vorname(self, orgs, board_a, api, people, koenig):
        org = orgs["a"]
        koenig.numbered = True
        koenig.save()
        otto = people["otto"].person
        Reign.objects.create(organization=org, dignity=koenig, year=1990, name="Otto Alt")
        Reign.objects.create(organization=org, dignity=koenig, year=2000, person=otto)
        Reign.objects.create(organization=org, dignity=koenig, year=2005, name="Erika Früh")
        # Otto Muster ein zweites Mal König: behält seine Zahl
        Reign.objects.create(organization=org, dignity=koenig, year=2010, person=otto)
        Reign.objects.create(organization=org, dignity=koenig, year=2020, name="Otto Neu")
        data = api(board_a).get(url(org, "reigns/")).json()
        names = {(r["year"], r["numbered_name"]) for r in data}
        assert names == {
            (1990, "Otto I. Alt"),
            (2000, "Otto II. Muster"),
            (2005, "Erika I. Früh"),
            (2010, "Otto II. Muster"),
            (2020, "Otto III. Neu"),
        }

    def test_jede_regentschaft_neu_zaehlen(self, orgs, board_a, api, people, koenig):
        """Ralf 2026-10-30: einstellbar, ob dieselbe Person bei jeder Würde neu gezählt wird."""
        org = orgs["a"]
        c = api(board_a)
        resp = c.patch(
            url(org, f"dignities/{koenig.pk}/"),
            {"numbered": True, "count_each_reign": True},
            "application/json",
        )
        assert resp.status_code == 200 and resp.json()["count_each_reign"] is True
        otto = people["otto"].person
        Reign.objects.create(organization=org, dignity=koenig, year=1990, name="Otto Alt")
        Reign.objects.create(organization=org, dignity=koenig, year=2000, person=otto)
        Reign.objects.create(organization=org, dignity=koenig, year=2010, person=otto)
        Reign.objects.create(
            organization=org, dignity=koenig, year=2015, name="Otto Neu", ordinal=7
        )
        Reign.objects.create(organization=org, dignity=koenig, year=2020, name="Otto Letzt")
        names = {r["year"]: r["numbered_name"] for r in c.get(url(org, "reigns/")).json()}
        assert names == {
            1990: "Otto I. Alt",
            2000: "Otto II. Muster",
            2010: "Otto III. Muster",
            2015: "Otto VII. Neu",
            2020: "Otto VIII. Letzt",
        }

    def test_zusatz_sen_jun_beim_nachnamen(self, orgs, board_a, api, koenig):
        org = orgs["a"]
        koenig.numbered = True
        koenig.save()
        Reign.objects.create(organization=org, dignity=koenig, year=1950, name="Josef Muster sen.")
        Reign.objects.create(organization=org, dignity=koenig, year=1980, name="Josef Muster jun.")
        data = api(board_a).get(url(org, "reigns/")).json()
        names = {r["year"]: r["numbered_name"] for r in data}
        assert names == {1950: "Josef I. Muster sen.", 1980: "Josef II. Muster jun."}

    def test_von_hand_und_ohne_schalter(self, orgs, board_a, api, koenig):
        org = orgs["a"]
        c = api(board_a)
        reigns = url(org, "reigns/")
        body = {"dignity": koenig.pk, "year": 1900, "name": "Michael Schmitz", "ordinal": 4}
        ok = c.post(reigns, body, "application/json")
        assert ok.status_code == 201, ok.content
        # ohne Schalter an der Würde: keine Zahl im Namen
        assert ok.json()["numbered_name"] == "Michael Schmitz"
        c.patch(url(org, f"dignities/{koenig.pk}/"), {"numbered": True}, "application/json")
        c.post(
            reigns,
            {"dignity": koenig.pk, "year": 1950, "name": "Michael Weber"},
            "application/json",
        )
        data = {r["year"]: r for r in c.get(reigns).json()}
        assert data[1900]["numbered_name"] == "Michael IV. Schmitz"
        assert data[1950]["numbered_name"] == "Michael V. Weber"  # zählt von IV weiter
        assert data[1950]["ordinal"] is None and data[1950]["ordinal_text"] == "V"
        bad = c.post(reigns, {**body, "ordinal": 0}, "application/json")
        assert bad.status_code == 400

    def test_oeffentliche_chronik(self, orgs, site, koenig):
        org = orgs["a"]
        koenig.numbered = True
        koenig.save()
        Reign.objects.create(organization=org, dignity=koenig, year=1950, name="Michael Schmitz")
        Reign.objects.create(organization=org, dignity=koenig, year=1960, name="Michael Weber")
        reigns = (
            Client()
            .get("/api/public/site/v1/schluessel-w/honors/")
            .json()["dignities"][0]["reigns"]
        )
        assert [x["holder"]["name"] for x in reigns] == ["Michael II. Weber", "Michael I. Schmitz"]


class TestMehrzahl:
    """Ralf 2026-10-30: nur eine Würde → ihr Name (Mehrzahl bei mehreren Regentschaften) im Menü."""

    def test_mehrzahl_raten(self):
        from honors.models import guess_plural

        assert [guess_plural(n) for n in ("Hubertuskönig", "Prinz", "Kaiser", "Königin")] == [
            "Hubertuskönige",
            "Prinzen",
            "Kaiser",
            "Königinnen",
        ]
        # Ehrungen (Ralf 2026-10-31)
        names = ("Ehrenmitglied", "Ehrenvorsitzender", "Ehrenmajor", "Ehrenhauptmann")
        assert [guess_plural(n) for n in names] == [
            "Ehrenmitglieder",
            "Ehrenvorsitzende",
            "Ehrenmajore",
            "Ehrenhauptleute",
        ]

    def test_bezeichnung_im_menue(self, orgs, site, koenig):
        org = orgs["a"]
        info = lambda: Client().get("/api/public/site/v1/schluessel-w/").json()  # noqa: E731
        assert info()["honors_label"] == "Majestäten"  # noch keine Regentschaft
        Reign.objects.create(organization=org, dignity=koenig, year=1950, name="Michael Schmitz")
        assert info()["honors_label"] == "König"
        Reign.objects.create(organization=org, dignity=koenig, year=1960, name="Michael Weber")
        assert info()["honors_label"] == "Könige"
        koenig.name_plural = "Schützenkönige"
        koenig.save()
        assert info()["honors_label"] == "Schützenkönige"
        honors = Client().get("/api/public/site/v1/schluessel-w/honors/").json()
        assert honors["dignities"][0]["name_plural"] == "Schützenkönige"
        prinz = Dignity.objects.create(organization=org, name="Prinz")
        Reign.objects.create(organization=org, dignity=prinz, year=1960, name="Paul Prinz")
        assert info()["honors_label"] == "Majestäten"
