"""Meldungen Schritt 3: der Verein schreibt für seine Untereinheiten aus, die Züge melden in
ihrem Bereich, der Verein bestätigt, fragt nach, erinnert; bestätigte Meldungen werden
Starter im eigenen Wettbewerb (Ralf 2026-11-01; alles erfunden)."""

import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from competitions.models import Competition, Score, Starter, Team
from members.models import DepartmentRoleHolder
from organizations.models import Department, DepartmentRole
from portal.models import PortalAccess

from .factories import give_role, make_member, make_user

pytestmark = pytest.mark.django_db
JSON = "application/json"


@pytest.fixture
def club(orgs):
    org = orgs["a"]
    org.enabled_modules = ["members", "competitions", "entries"]
    org.department_access_enabled = True
    org.member_portal_enabled = True
    org.save()
    return org


@pytest.fixture
def zuege(club):
    return {
        name: Department.objects.create(
            self_managed=True, organization=club, name=name, short_name=short
        )
        for name, short in (("Jägerzug", "Jäger"), ("Grenadiere", "Gren"), ("Tellschützen", "Tell"))
    }


@pytest.fixture
def leute(club, zuege):
    j, g = zuege["Jägerzug"], zuege["Grenadiere"]
    return {
        "Max": make_member(club, "Max", "Muster", departments=[j]),
        "Erika": make_member(club, "Erika", "Beispiel", departments=[j]),
        "Otto": make_member(club, "Otto", "Normal", departments=[g]),
        "Paul": make_member(club, "Paul", "Weber", departments=[g]),
    }


def zugfuehrer(club, zug, membership, email, perms=("dept.entries.submit", "dept.members.view")):
    role = DepartmentRole.objects.create(department=zug, name="Zugführer", permissions=list(perms))
    DepartmentRoleHolder.objects.create(
        organization=club, membership=membership, role=role, valid_from=datetime.date(2020, 1, 1)
    )
    user = make_user(email)
    PortalAccess.objects.create(organization=club, membership=membership, user=user)
    return user


def org_url(club, rest=""):
    return f"/api/orgs/{club.pk}/entries/{rest}"


def dept_url(zug, rest=""):
    return f"/api/depts/{zug.pk}/entries/{rest}"


def _ok(resp, code=(200, 201)):
    assert resp.status_code in ((code,) if isinstance(code, int) else code), resp.content
    return resp.json()


def _ausschreibung(c, club, zuege, **extra):
    data = {
        "title": "Preisschießen 2027",
        "internal": True,
        "date": "2027-05-09",
        "deadline": "2027-04-30",
        "location": "Vereinsheim",
        "template": "mannschaften",
        "units": [zuege["Jägerzug"].pk, zuege["Grenadiere"].pk],
        **extra,
    }
    call = _ok(c.post(org_url(club, "calls/"), data, JSON))
    block = call["blocks"][0]["id"]
    _ok(
        c.patch(
            org_url(club, f"calls/{call['id']}/blocks/{block}/"),
            {"settings": {"min_teams": 1, "max_teams": 1, "min_size": 2, "max_size": 2}},
            JSON,
        )
    )
    return call


class TestAusschreibung:
    def test_anlegen_nur_mit_recht(self, club, zuege, api):
        melder = make_user("melder@example.org")
        give_role(melder, club, ["entries.submit"], name="Melder")
        resp = api(melder).post(
            org_url(club, "calls/"), {"title": "Preisschießen", "internal": True}, JSON
        )
        assert resp.status_code == 403

    def test_veranstalter_und_untereinheiten(self, club, zuege, orgs, board_a, api):
        c = api(board_a)
        call = _ausschreibung(c, club, zuege)
        assert call["internal"] and call["organizer"] == club.name
        # Die Einheit selbst meldet nicht
        resp = c.post(org_url(club, f"calls/{call['id']}/entry/"))
        assert resp.status_code == 400
        rows = _ok(c.get(org_url(club, f"calls/{call['id']}/units/")))
        assert [r["name"] for r in rows] == ["Jägerzug", "Grenadiere"]
        assert all(r["status"] is None for r in rows)
        listed = _ok(c.get(org_url(club, "calls/")))
        assert listed[0]["unit_counts"] == {"addressed": 2, "submitted": 0}
        # Ohne Auswahl: alle direkten Untereinheiten
        _ok(c.patch(org_url(club, f"calls/{call['id']}/"), {"units": []}, JSON))
        rows = _ok(c.get(org_url(club, f"calls/{call['id']}/units/")))
        assert len(rows) == 3
        # Fremde Untereinheit
        fremd = Department.objects.create(organization=orgs["b"], name="Fremd")
        resp = c.patch(org_url(club, f"calls/{call['id']}/"), {"units": [fremd.pk]}, JSON)
        assert resp.status_code == 400


class TestAblauf:
    def test_zug_meldet_verein_bestaetigt(self, club, zuege, leute, board_a, api):
        c = api(board_a)
        call = _ausschreibung(c, club, zuege)
        jaeger = zuege["Jägerzug"]
        chef = api(zugfuehrer(club, jaeger, leute["Max"], "max@example.org"))
        # Der Zug sieht die Ausschreibung in seinem Bereich
        calls = _ok(chef.get(dept_url(jaeger, "calls/")))
        assert [x["title"] for x in calls] == ["Preisschießen 2027"]
        base = dept_url(jaeger, f"calls/{call['id']}/")
        # … darf sie aber nicht ändern
        assert chef.patch(base, {"title": "X"}, JSON).status_code == 403
        block = call["blocks"][0]["id"]
        entry = _ok(chef.post(f"{base}entry/teams/", {"block": block}))
        team = entry["teams"][0]
        assert team["name"] == "Jäger I"
        # Auswahl nur aus dem eigenen Zug
        names = [r["name"] for r in _ok(chef.get(f"{base}candidates/?suche=o"))]
        assert "Otto Normal" not in names
        names = [r["name"] for r in _ok(chef.get(f"{base}candidates/?suche=Muster"))]
        assert names == ["Max Muster"]
        resp = chef.post(
            f"{base}entry/persons/",
            {"block": block, "team": team["id"], "membership": leute["Otto"].pk},
            JSON,
        )
        assert resp.status_code == 400
        for first in ("Max", "Erika"):
            _ok(
                chef.post(
                    f"{base}entry/persons/",
                    {"block": block, "team": team["id"], "membership": leute[first].pk},
                    JSON,
                )
            )
        entry = _ok(chef.post(f"{base}entry/submit/", {"channel": "brief"}, JSON))
        assert entry["status"] == "abgegeben"
        assert entry["versions"][0]["channel"] == "hubiplan"
        assert entry["unit_name"] == "Jägerzug"
        # Bestätigen darf nur der Verein
        assert chef.post(f"{base}entry/confirm/").status_code == 403

        rows = {r["name"]: r for r in _ok(c.get(org_url(club, f"calls/{call['id']}/units/")))}
        assert rows["Jägerzug"]["status"] == "abgegeben"
        assert rows["Jägerzug"]["persons"] == 2
        assert rows["Grenadiere"]["status"] is None
        unit = org_url(club, f"units/{jaeger.pk}/calls/{call['id']}/")
        entry = _ok(c.post(f"{unit}entry/confirm/", {"note": ""}, JSON))
        assert entry["status"] == "bestaetigt"
        comp = Competition.objects.get(pk=entry["competition"])
        assert comp.own and comp.title == "Preisschießen 2027"
        starters = Starter.objects.filter(competition=comp)
        assert sorted(s.name for s in starters) == ["Erika Beispiel", "Max Muster"]
        assert {s.unit_id for s in starters} == {jaeger.pk}
        assert [t.name for t in Team.objects.filter(competition=comp)] == ["Jäger I"]
        assert sorted(s.name for s in comp.scorings.all()) == [
            "Einzelwertung",
            "Mannschaftswertung",
            "Wertung der Untereinheiten",
        ]

        # Ersatz nach Bestätigung, erneut abgegeben und bestätigt: Starter folgt
        person = next(p for p in entry["persons"] if p["name"] == "Erika Beispiel")
        willi = make_member(club, "Willi", "Wacker", departments=[jaeger])
        _ok(chef.patch(f"{base}entry/persons/{person['id']}/", {"membership": willi.pk}, JSON))
        _ok(chef.post(f"{base}entry/submit/", {}, JSON))
        _ok(c.post(f"{unit}entry/confirm/", {}, JSON))
        assert sorted(s.name for s in Starter.objects.filter(competition=comp)) == [
            "Max Muster",
            "Willi Wacker",
        ]
        # Ein Starter, den der Verein selbst eingetragen hat, bleibt unberührt
        Starter.objects.create(organization=club, competition=comp, name="Gerd Gast")
        _ok(chef.post(f"{base}entry/submit/", {"withdraw": True, "reason": "krank"}, JSON))
        assert sorted(s.name for s in Starter.objects.filter(competition=comp)) == ["Gerd Gast"]

    def test_rueckfrage_und_erinnerung(self, club, zuege, leute, board_a, api, mailoutbox):
        c = api(board_a)
        call = _ausschreibung(c, club, zuege)
        jaeger, gren = zuege["Jägerzug"], zuege["Grenadiere"]
        chef = api(zugfuehrer(club, jaeger, leute["Max"], "max@example.org"))
        zugfuehrer(club, gren, leute["Otto"], "otto@example.org")
        base = dept_url(jaeger, f"calls/{call['id']}/")
        block = call["blocks"][0]["id"]
        team = _ok(chef.post(f"{base}entry/teams/", {"block": block}))["teams"][0]["id"]
        _ok(
            chef.post(
                f"{base}entry/persons/",
                {"block": block, "team": team, "membership": leute["Max"].pk},
                JSON,
            )
        )
        _ok(chef.post(f"{base}entry/submit/", {}, JSON))

        # Erinnerung nur an die Grenadiere (Jäger haben abgegeben)
        mailoutbox.clear()
        with self._commit():
            res = _ok(c.post(org_url(club, f"calls/{call['id']}/remind/"), {}, JSON))
        assert res == {"reminded": ["Grenadiere"], "without_contact": []}
        assert [m.to for m in mailoutbox] == [["otto@example.org"]]
        assert "Preisschießen 2027" in mailoutbox[0].subject

        # Rückfrage an die Jäger
        unit = org_url(club, f"units/{jaeger.pk}/calls/{call['id']}/")
        mailoutbox.clear()
        with self._commit():
            entry = _ok(c.post(f"{unit}entry/return/", {"question": "Zweiter Schütze fehlt"}, JSON))
        assert entry["status"] == "rueckfrage" and entry["question"] == "Zweiter Schütze fehlt"
        assert [m.to for m in mailoutbox] == [["max@example.org"]]
        seen = _ok(chef.get(f"{base}entry/"))
        assert seen["question"] == "Zweiter Schütze fehlt"
        # Ohne Änderung erneut abgeben: beantwortet, ohne neue Fassung
        entry = _ok(chef.post(f"{base}entry/submit/", {}, JSON))
        assert entry["status"] == "abgegeben" and entry["question"] == ""
        assert len(entry["versions"]) == 1
        # Die Zugführung darf nicht zurückgeben
        assert chef.post(f"{base}entry/return/", {"question": "x"}, JSON).status_code == 403

    @staticmethod
    def _commit():
        from django.test import TestCase

        return TestCase.captureOnCommitCallbacks(execute=True)

    def test_strenge_bedingungen_und_doppelt_gemeldet(self, club, zuege, leute, board_a, api):
        c = api(board_a)
        call = _ausschreibung(c, club, zuege, strict=True)
        jaeger, gren = zuege["Jägerzug"], zuege["Grenadiere"]
        block = call["blocks"][0]["id"]
        # Der Verein trägt für beide Züge ein (von oben verwaltet), auch Mitglieder anderer Züge
        bases = {z: org_url(club, f"units/{z.pk}/calls/{call['id']}/") for z in (jaeger, gren)}
        teams = {
            z: _ok(c.post(f"{b}entry/teams/", {"block": block}))["teams"][0]["id"]
            for z, b in bases.items()
        }
        for z, firsts in ((jaeger, ("Max",)), (gren, ("Otto", "Max"))):
            for first in firsts:
                _ok(
                    c.post(
                        f"{bases[z]}entry/persons/",
                        {"block": block, "team": teams[z], "membership": leute[first].pk},
                        JSON,
                    )
                )
        entry = _ok(c.get(f"{bases[gren]}entry/"))
        assert any("auch von Jägerzug gemeldet" in h["text"] for h in entry["hints"])
        resp = c.post(f"{bases[jaeger]}entry/submit/", {}, JSON)
        assert resp.status_code == 400
        assert "1 von mindestens 2" in str(resp.json()["hints"])

    def test_meldeliste(self, club, zuege, leute, board_a, api):
        c = api(board_a)
        call = _ausschreibung(c, club, zuege)
        jaeger = zuege["Jägerzug"]
        base = org_url(club, f"units/{jaeger.pk}/calls/{call['id']}/")
        block = call["blocks"][0]["id"]
        team = _ok(c.post(f"{base}entry/teams/", {"block": block}))["teams"][0]["id"]
        for first in ("Max", "Erika"):
            _ok(
                c.post(
                    f"{base}entry/persons/",
                    {"block": block, "team": team, "membership": leute[first].pk},
                    JSON,
                )
            )
        _ok(c.post(f"{base}entry/submit/", {}, JSON))
        resp = c.get(org_url(club, f"calls/{call['id']}/list/?als=csv"))
        assert resp.status_code == 200
        text = resp.content.decode("utf-8-sig")
        assert "Jägerzug;Mannschaften;Jäger I;1;Max Muster" in text
        resp = c.get(org_url(club, f"calls/{call['id']}/list/?als=xlsx"))
        import io

        ws = load_workbook(io.BytesIO(resp.content)).active
        assert ws["A1"].value == "Untereinheit" and ws.max_row == 3


class TestZugang:
    def test_ohne_abteilungsrecht(self, club, zuege, leute, board_a, api):
        c = api(board_a)
        call = _ausschreibung(c, club, zuege)
        jaeger = zuege["Jägerzug"]
        nur_lesen = api(
            zugfuehrer(club, jaeger, leute["Max"], "max@example.org", perms=("dept.members.view",))
        )
        assert nur_lesen.get(dept_url(jaeger, "calls/")).status_code == 403
        # Nicht angeschriebene Untereinheit
        tell = zuege["Tellschützen"]
        chef = api(zugfuehrer(club, tell, leute["Paul"], "paul@example.org"))
        assert _ok(chef.get(dept_url(tell, "calls/"))) == []
        assert chef.get(dept_url(tell, f"calls/{call['id']}/")).status_code == 404

    def test_melder_ohne_verwaltung_sieht_interne_nicht(self, club, zuege, board_a, api):
        _ausschreibung(api(board_a), club, zuege)
        melder = make_user("melder@example.org")
        give_role(melder, club, ["entries.submit"], name="Melder")
        assert _ok(api(melder).get(org_url(club, "calls/"))) == []

    def test_modul_aus(self, club, zuege, leute, api):
        jaeger = zuege["Jägerzug"]
        chef = api(zugfuehrer(club, jaeger, leute["Max"], "max@example.org"))
        club.enabled_modules = ["members"]
        club.save()
        assert chef.get(dept_url(jaeger, "calls/")).status_code == 404


# --- Wettbewerbe Schritt 3, Rest: Ergebnisse zurück an die meldenden Untereinheiten ---
# Eigene Starter, Mannschaften und die Wertung der Untereinheiten immer; die ganze
# Einzelrangliste nur mit Freigabe des Ausrichters (Konzept Wettbewerbe Abschnitt 9).


def _melden(c, chef, club, zug, call, leute_):
    """Zug meldet (ohne ``chef``: der Verein trägt die Meldung für ihn ein), Verein bestätigt."""
    unit = org_url(club, f"units/{zug.pk}/calls/{call['id']}/")
    base = dept_url(zug, f"calls/{call['id']}/") if chef else unit
    chef = chef or c
    block = call["blocks"][0]["id"]
    entry = _ok(chef.post(f"{base}entry/teams/", {"block": block}))
    team = entry["teams"][0]["id"]
    for m in leute_:
        _ok(
            chef.post(
                f"{base}entry/persons/",
                {"block": block, "team": team, "membership": m.pk},
                JSON,
            )
        )
    _ok(chef.post(f"{base}entry/submit/", {"channel": "brief"}, JSON))
    return _ok(c.post(f"{unit}entry/confirm/", {"note": ""}, JSON)), base


def test_ergebnisse_zurueck(club, zuege, leute, board_a, api):
    c = api(board_a)
    call = _ausschreibung(c, club, zuege)
    jaeger, gren = zuege["Jägerzug"], zuege["Grenadiere"]
    chef = api(zugfuehrer(club, jaeger, leute["Max"], "max@example.org"))
    zugfuehrer(club, gren, leute["Otto"], "otto@example.org")
    entry, base = _melden(c, chef, club, jaeger, call, [leute["Max"], leute["Erika"]])
    _melden(c, None, club, gren, call, [leute["Otto"], leute["Paul"]])
    comp = Competition.objects.get(pk=entry["competition"])
    einzel = comp.scorings.get(kind="einzel")
    zugw = comp.scorings.get(kind="einheit")
    zugw.source = einzel
    zugw.team_size = 2
    zugw.save()
    for name, value in (("Max", 30), ("Erika", 25), ("Otto", 28), ("Paul", 27)):
        s = Starter.objects.get(competition=comp, membership=leute[name])
        Score.objects.create(organization=club, scoring=einzel, starter=s, value=Decimal(value))

    # Vor dem Abschluss: nur der Zustand
    data = _ok(chef.get(f"{base}entry/results/"))
    assert data["final"] is False and data["scorings"] == []
    comp.status = Competition.Status.FINAL
    comp.save()

    data = _ok(chef.get(f"{base}entry/results/"))
    by = {s["name"]: s for s in data["scorings"]}
    # Einzel: nur eigene Starter, mit Platz in der Gesamtrangliste
    rows = by["Einzelwertung"]["rows"]
    assert [(r["place"], r["name"]) for r in rows] == [(1, "Max Muster"), (4, "Erika Beispiel")]
    assert all(r["own"] for r in rows) and not by["Einzelwertung"]["complete"]
    assert by["Einzelwertung"]["field"] == 4
    # Wertung der Untereinheiten: alle Züge, der eigene markiert
    zug_rows = by["Wertung der Untereinheiten"]["rows"]
    assert [(r["name"], r["own"]) for r in zug_rows] == [
        ("Grenadiere", False),
        ("Jägerzug", True),
    ]
    # Freigabe: ganze Rangliste mit fremden Namen
    comp.share_ranking = True
    comp.save()
    rows = _ok(chef.get(f"{base}entry/results/"))["scorings"][0]["rows"]
    assert [r["name"] for r in rows] == [
        "Max Muster",
        "Otto Normal",
        "Paul Weber",
        "Erika Beispiel",
    ]
    assert [r["own"] for r in rows] == [True, False, False, True]


def test_ohne_wettbewerb_oder_fremder_zug(club, zuege, leute, board_a, api):
    c = api(board_a)
    call = _ausschreibung(c, club, zuege, create_competition=False)
    jaeger = zuege["Jägerzug"]
    chef = api(zugfuehrer(club, jaeger, leute["Max"], "max@example.org"))
    base = dept_url(jaeger, f"calls/{call['id']}/")
    # Keine Meldung, kein Wettbewerb: 404
    assert chef.get(f"{base}entry/results/").status_code == 404
    # Andere Untereinheit sieht nichts von dieser Meldung
    tell = zuege["Tellschützen"]
    other = api(
        zugfuehrer(
            club, tell, make_member(club, "Tim", "Tell", departments=[tell]), "tim@example.org"
        )
    )
    assert other.get(dept_url(tell, f"calls/{call['id']}/entry/results/")).status_code == 404
