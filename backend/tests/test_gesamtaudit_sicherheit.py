"""Gesamtaudit 25.10.2026, Paket 1: Sicherheitsgrenzen (Befunde A-, B-).

Die Prüfproben des Audits zeigten das fehlerhafte Verhalten; hier steht das
gewünschte Verhalten. Alles mit erfundenen Daten, ohne echten Mailversand.
"""

import datetime
import io
import json
import zipfile
from decimal import Decimal

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.mail import EmailMessage
from django.http import FileResponse
from django.test import Client, RequestFactory, override_settings
from openpyxl import Workbook, load_workbook
from PIL import Image

from billing import services as billing_services
from billing.models import BillingSettings, Invoice, InvoiceLine, Payment
from core.filecheck import SafeFileMiddleware, sniff
from core.tenancy import rls_bypass, tenant_context
from exchange import reading
from finance.models import Account, Category
from organizations.models import Department
from outbox.backend import OutboxBackend
from outbox.models import OutgoingMail

from .factories import PASSWORD, give_role, make_member, make_user

pytestmark = pytest.mark.django_db
JSON = "application/json"


@pytest.fixture(autouse=True)
def _isolated(settings, tmp_path, monkeypatch):
    settings.PRIVATE_ROOT = tmp_path / "private"
    settings.CELERY_TASK_ALWAYS_EAGER = True
    monkeypatch.setattr("django.db.transaction.on_commit", lambda func, *a, **k: func())


def png_bytes(size=(12, 8)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGBA", size, (200, 10, 10, 255)).save(buf, "PNG")
    return buf.getvalue()


def department_doc():
    return {
        "format": "hubiplan-export",
        "version": 1,
        "kind": "department",
        "persons": [],
        "member_types": [],
        "memberships": [],
        "department": {"id": 1, "name": "Prüfzug", "founding_year": 2000},
    }


def start_import(client, club, doc, files=None):
    base = f"/api/orgs/{club.pk}/datatransfer/imports/"
    raw = json.dumps(doc).encode()
    name = "export.json"
    if files:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("export.json", raw)
            for path, content in files.items():
                archive.writestr(path, content)
        raw, name = buf.getvalue(), "export.zip"
    resp = client.post(base, {"file": SimpleUploadedFile(name, raw)})
    assert resp.status_code == 201, resp.content
    job_id = resp.json()["id"]
    return base, job_id, client.get(f"{base}{job_id}/").json()


# --- B-S02: Abteilungsimport schreibt keine Vereinskasse ---


def test_abteilungsimport_lehnt_finanztabellen_ab(api, orgs):
    club = orgs["c"]
    club.enabled_modules = ["members", "finance"]
    club.save()
    user = make_user("nur-org@example.org")
    give_role(user, club, permissions=["org.edit"], name="Nur Organisation")
    client = api(user)
    doc = department_doc()
    doc["tables"] = {
        "finance.account": [
            {
                "id": 17,
                "name": "Eingeschleust",
                "opening_date": "2025-01-01",
                "opening_balance": "9999.00",
            }
        ]
    }
    base, job_id, job = start_import(client, club, doc)
    assert job["preview"]["error_count"] >= 1
    assert any("nicht übernommen" in e for e in job["preview"]["errors"])
    assert client.post(f"{base}{job_id}/run/", {}, content_type=JSON).status_code == 400
    with tenant_context(club.pk):
        assert not Account.objects.filter(name="Eingeschleust").exists()
        assert not Department.objects.filter(name="Prüfzug", organization=club).exists()


def test_abteilungsimport_ohne_kasse_geht_weiter(api, orgs):
    club = orgs["c"]
    user = make_user("nur-org2@example.org")
    give_role(user, club, permissions=["org.edit"], name="Nur Organisation")
    client = api(user)
    base, job_id, job = start_import(client, club, department_doc())
    assert job["preview"]["error_count"] == 0, job["preview"]
    assert client.post(f"{base}{job_id}/run/", {}, content_type=JSON).status_code == 200
    with tenant_context(club.pk):
        assert Department.objects.filter(name="Prüfzug", organization=club).exists()


# --- B-S01: Dateien aus dem Import werden wie beim Hochladen geprüft ---


def test_html_als_logo_wird_abgewiesen(api, orgs, settings):
    club = orgs["c"]
    user = make_user("logo@example.org")
    give_role(user, club, permissions=["org.edit"], name="Nur Organisation")
    client = api(user)
    doc = department_doc()
    doc["with_files"] = True
    doc["department"]["logo"] = "files/logo-probe.html"
    harmless = b"<!doctype html><script>document.title='PROBE'</script>"
    base, job_id, _ = start_import(client, club, doc, {"files/logo-probe.html": harmless})
    client.post(f"{base}{job_id}/run/", {}, content_type=JSON)
    assert client.get(f"{base}{job_id}/").json()["status"] != "done"
    with tenant_context(club.pk):
        assert not Department.objects.filter(name="Prüfzug", organization=club).exists()
    media = settings.MEDIA_ROOT
    assert not list(media.rglob("*.html")) if media.exists() else True


def test_echtes_logo_wird_als_png_neu_geschrieben(api, orgs):
    club = orgs["c"]
    user = make_user("logo2@example.org")
    give_role(user, club, permissions=["org.edit"], name="Nur Organisation")
    client = api(user)
    doc = department_doc()
    doc["with_files"] = True
    doc["department"]["logo"] = "files/logo.html"  # Endung egal: Inhalt zählt
    base, job_id, _ = start_import(client, club, doc, {"files/logo.html": png_bytes()})
    client.post(f"{base}{job_id}/run/", {}, content_type=JSON)
    assert client.get(f"{base}{job_id}/").json()["status"] == "done"
    with tenant_context(club.pk):
        dept = Department.objects.get(name="Prüfzug", organization=club)
        assert dept.logo.name.endswith(".png")
        with dept.logo.open("rb") as handle:
            assert sniff(handle.read()) == "image/png"


def test_logo_upload_wird_als_png_neu_geschrieben(api, orgs, board_a):
    buf = io.BytesIO()
    Image.new("RGB", (20, 10), "blue").save(buf, "JPEG")
    buf.seek(0)
    buf.name = "logo.jpg"
    resp = api(board_a).post(f"/api/orgs/{orgs['a'].pk}/logo/", {"logo": buf})
    assert resp.status_code == 200, resp.content
    assert resp.json()["logo_url"].endswith(".png")
    orgs["a"].refresh_from_db()
    with orgs["a"].logo.open("rb") as handle:
        assert sniff(handle.read()) == "image/png"


def test_importierter_beleg_typ_kommt_aus_dem_inhalt():
    from datatransfer.tables import check_file
    from finance.models import Receipt

    receipt = Receipt(content_type="text/html")
    with pytest.raises(ValidationError):
        check_file("finance.receipt", receipt, b"<html><script></script></html>", "files/x")
    check_file("finance.receipt", receipt, b"%PDF-1.4 ...", "files/x")
    assert receipt.content_type == "application/pdf"


def test_dateien_nie_als_aktive_seite_ausgeliefert(tmp_path):
    path = tmp_path / "beleg"
    path.write_bytes(b"<script>alert(1)</script>")
    middleware = SafeFileMiddleware(
        lambda request: FileResponse(path.open("rb"), content_type="text/html", filename="b.html")
    )
    response = middleware(RequestFactory().get("/api/x/"))
    assert response["Content-Disposition"].startswith("attachment")
    assert "sandbox" in response["Content-Security-Policy"]
    response.close()
    pdf = SafeFileMiddleware(
        lambda request: FileResponse(
            io.BytesIO(b"%PDF-1.4"), content_type="application/pdf", filename="a.pdf"
        )
    )(RequestFactory().get("/api/x/"))
    assert pdf["Content-Disposition"].startswith("inline")


# --- B-S05: Grenzen greifen früh ---


def test_excel_bricht_frueh_ab(tmp_path, monkeypatch):
    monkeypatch.setattr(reading, "MAX_ROWS", 20)
    monkeypatch.setattr(reading, "MAX_SCAN_ROWS", 30)
    path = tmp_path / "gross.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.append(["name", "wert"])
    for index in range(2000):
        sheet.append([f"zeile-{index}", index])
    workbook.save(path)
    seen = []
    original = reading._cell

    def counting(value):
        seen.append(value)
        return original(value)

    monkeypatch.setattr(reading, "_cell", counting)
    with pytest.raises(reading.ReadError, match="20"):
        reading.read_table(path, path.name)
    assert len(seen) < 200  # nicht 4000 Zellen gelesen


def test_zip_anhang_ueber_grenze(monkeypatch):
    from datatransfer import load

    monkeypatch.setattr(load, "MAX_FILE_BYTES", 100)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("files/gross", b"0" * 10_000)  # stark komprimierbar
    files = load.Files(zipfile.ZipFile(buf))
    with pytest.raises(load.FileProblem):
        files.read("files/gross")


# --- B-S03: veralteter Rechnungsstand ---


def test_zweite_vollzahlung_auf_altem_stand_wird_abgelehnt(orgs, board_c):
    club = orgs["c"]
    member = make_member(club, "Paul", "Zahler")
    with tenant_context(club.pk):
        account = Account.objects.create(
            organization=club, name="Bank", opening_date=datetime.date(2025, 1, 1)
        )
        category = Category.objects.create(organization=club, name="Beiträge", kind="income")
        BillingSettings.objects.create(organization=club, account=account, category=category)
        invoice = Invoice.objects.create(
            organization=club,
            membership=member,
            recipient_kind="member",
            recipient_name="Paul Zahler",
            date=datetime.date(2025, 1, 1),
            due_date=datetime.date(2025, 1, 15),
            year=2025,
            number=1,
            status="open",
            total=Decimal("100.00"),
        )
        InvoiceLine.objects.create(organization=club, invoice=invoice, text="Beitrag", amount=100)
        first = Invoice.objects.get(pk=invoice.pk)
        stale = Invoice.objects.get(pk=invoice.pk)
        kwargs = {
            "date": datetime.date(2025, 2, 1),
            "amount": Decimal("100.00"),
            "method": "transfer",
            "account": account,
            "user": board_c,
        }
        billing_services.add_payment(first, **kwargs)
        with pytest.raises(ValidationError):
            billing_services.add_payment(stale, **kwargs)
        invoice.refresh_from_db()
        assert invoice.paid == Decimal("100.00")
        assert Payment.objects.filter(invoice=invoice).count() == 1


# --- B-S04: Formeln in der Teilnehmerliste ---


def test_bemerkung_bleibt_text_in_teilnehmerliste(orgs):
    from django.utils import timezone

    from events.models import Event, EventResponse
    from events.participants import Row, xlsx

    start = timezone.make_aware(datetime.datetime(2027, 1, 4, 10))
    event = Event(
        pk=1,
        organization=orgs["c"],
        title="=HYPERLINK(1)",
        starts_at=start,
        ends_at=start + datetime.timedelta(hours=1),
    )
    response = EventResponse(status="yes", guests=2, note="=1+1")
    row = Row(1, "Erika Muster", "42", "1. Zug", response)
    with tenant_context(orgs["c"].pk):
        content = xlsx(event, [row], [])
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    ws = workbook.active
    assert ws["F4"].data_type == "s"
    assert ws["F4"].value == "'=1+1"
    assert ws["A1"].data_type == "s"
    assert ws["E4"].value == 2  # Zahlen bleiben Zahlen
    workbook.close()


# --- B-F01: Mailserver nimmt nicht an ---


def test_smtp_null_ist_kein_versand(settings, monkeypatch):
    settings.EMAIL_HOST = "unused.invalid"

    class ZeroDelivery:
        def __init__(self, **kwargs):
            pass

        def send_messages(self, messages):
            return 0

    monkeypatch.setattr("outbox.backend.SmtpBackend", ZeroDelivery)
    backend = OutboxBackend(fail_silently=True)
    message = EmailMessage("Probe", "nur lokal", "probe@example.org", ["test@example.org"])
    assert backend.send_messages([message]) == 0
    with rls_bypass():
        assert OutgoingMail.objects.get(subject="Probe").status == OutgoingMail.Status.FAILED


@override_settings(EMAIL_BACKEND="outbox.backend.OutboxBackend", EMAIL_HOST="unused.invalid")
def test_mail_send_meldet_fehler(monkeypatch):
    from outbox import mail

    class ZeroDelivery:
        def __init__(self, **kwargs):
            pass

        def send_messages(self, messages):
            return 0

    monkeypatch.setattr("outbox.backend.SmtpBackend", ZeroDelivery)
    assert mail.send("x@example.org", "Probe 2", "Text", fail_silently=True) is False


# --- A-S01: keine Anmeldegeheimnisse im Postausgang ---


@override_settings(EMAIL_BACKEND="outbox.backend.OutboxBackend")
def test_postausgang_speichert_keine_anmeldelinks(settings):
    from outbox import mail

    body = (
        f"Link: {settings.FRONTEND_URL}/anmelden/link/GEHEIM123\n"
        f"Passwort: {settings.FRONTEND_URL}/passwort-zuruecksetzen/abc-GEHEIM456\n"
        f"E-Mail: {settings.FRONTEND_URL}/konto/email/GEHEIM789\n"
        f"Termine: {settings.FRONTEND_URL}/termine\n"
    )
    mail.send("x@example.org", "Zugang", body)
    with rls_bypass():
        stored = OutgoingMail.objects.get(subject="Zugang").body
    assert "GEHEIM" not in stored
    assert f"{settings.FRONTEND_URL}/termine" in stored


# --- A-S04 / A-S05: Anmeldelinks ---


def test_passwortwechsel_widerruft_offene_links(api):
    from portal.models import LoginLink
    from portal.services import create_login_link

    user = make_user("wechsel@example.org")
    token = create_login_link(user).rsplit("/", 1)[-1]
    resp = api(user).post(
        "/api/me/password/",
        {"current_password": PASSWORD, "new_password": "Ganz-Neues-Passwort-42"},
        content_type=JSON,
    )
    assert resp.status_code == 200, resp.content
    with rls_bypass():
        assert not LoginLink.objects.get(user=user).is_open
    assert Client().post(f"/api/public/login-link/{token}/").status_code == 404
    # neue Links gehen weiter
    fresh = create_login_link(user).rsplit("/", 1)[-1]
    assert Client().post(f"/api/public/login-link/{fresh}/").status_code == 200


def test_anmeldelink_nur_mit_csrf(client):
    from portal.services import create_login_link

    user = make_user("csrf@example.org")
    token = create_login_link(user).rsplit("/", 1)[-1]
    # fremde Seite ohne CSRF-Token: abgelehnt, Link bleibt gültig
    assert client.post(f"/api/public/login-link/{token}/").status_code == 403
    from .conftest import csrf_headers

    resp = client.post(f"/api/public/login-link/{token}/", headers=csrf_headers(client))
    assert resp.status_code == 200, resp.content


# --- A-S02: Rückgängig nur mit dem Recht für die Felder ---


def test_gestaltungsrecht_nimmt_keine_einstellungen_zurueck(api, orgs, board_a):
    org = orgs["a"]
    resp = api(board_a).patch(
        f"/api/orgs/{org.pk}/", {"contact_changes_need_approval": True}, content_type=JSON
    )
    assert resp.status_code == 200, resp.content
    change_set = resp["X-Change-Set"]
    designer = make_user("gestalter@example.org")
    give_role(designer, org, permissions=["theme.edit", "audit.view"], name="Gestaltung")
    undo = api(designer).post(
        f"/api/orgs/{org.pk}/changes/{change_set}/undo/", {}, content_type=JSON
    )
    assert undo.status_code == 403

    resp = api(board_a).patch(
        f"/api/orgs/{org.pk}/theme/", {"primary_color": "#123E7A"}, content_type=JSON
    )
    assert resp.status_code == 200
    undo = api(designer).post(
        f"/api/orgs/{org.pk}/changes/{resp['X-Change-Set']}/undo/", {}, content_type=JSON
    )
    assert undo.status_code == 200, undo.content


# --- A-F01: eigene Rollenverwaltung nicht über Umwege beenden ---


def test_eigene_rollenverwaltung_nicht_per_datum_oder_rolle_entziehen(api, orgs):
    org = orgs["a"]
    admin = make_user("allein@example.org")
    assignment = give_role(admin, org, permissions=["roles.manage", "org.edit"], name="Admin")
    c = api(admin)
    yesterday = (datetime.date.today() - datetime.timedelta(days=1)).isoformat()
    resp = c.patch(
        f"/api/orgs/{org.pk}/assignments/{assignment.pk}/",
        {"valid_until": yesterday},
        content_type=JSON,
    )
    assert resp.status_code == 400, resp.content
    resp = c.patch(
        f"/api/orgs/{org.pk}/roles/{assignment.role_id}/",
        {"permissions": ["org.edit"]},
        content_type=JSON,
    )
    assert resp.status_code == 400, resp.content
    assignment.refresh_from_db()
    assert assignment.valid_until is None
    assert "roles.manage" in assignment.role.permissions
